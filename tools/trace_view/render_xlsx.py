# -*- coding: utf-8 -*-
"""Excel workbook: sessions, steps, gaps, findings and per-tool statistics, with colouring."""
from __future__ import annotations

import re

from . import labels as L
from .render_md import merged_rows, status_text  # noqa: F401  (status_text re-used below)
from .timeutil import fmt_dt, fmt_duration, fmt_time

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe(value):
    """Strings only ever land as text: no control characters, and nothing that starts a formula."""
    if isinstance(value, str):
        value = _ILLEGAL.sub("", value)
        if value[:1] in ("=", "+", "-", "@"):
            value = "'" + value
    return value


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append([_safe(h) for h in header])
    for r in rows:
        ws.append([_safe(c) for c in r])
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE3EA")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (_col(len(header)), len(rows) + 1)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[_col(i)].width = w
    return ws


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def render_xlsx(report, tz, path, lang="ja", max_steps=None) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
    from openpyxl.styles import PatternFill

    lab = L.get(lang)
    wb = Workbook()
    wb.remove(wb.active)

    rows = []
    for s in report["sessions"]:
        st = s["stats"]["status"]
        rows.append([s["id"] or "(none)", fmt_dt(s["first"], tz), fmt_dt(s["last"], tz), len(s["steps"]),
                     st.get("ok", 0), st.get("error", 0), st.get("refused", 0), st.get("unavailable", 0),
                     st.get("no_outcome", 0), len(s["gaps"]), lab[s["concurrency"]["mode"]],
                     s["concurrency"]["max_overlap"], len(s["findings"]),
                     ", ".join(s["connect"]["methods"])])
    ws = _sheet(wb, lab["sessions"], [lab["session"], lab["first"], lab["last"], lab["calls"], lab["ok"],
                                      lab["error"], lab["refused"], lab["unavailable"], lab["no_outcome"],
                                      lab["gaps"], lab["mode"], lab["max_overlap"], lab["findings"],
                                      lab["connect"]],
                rows, [20, 20, 20, 9, 8, 8, 8, 10, 10, 8, 10, 10, 10, 40])
    if rows:
        n = len(rows) + 1
        ws.conditional_formatting.add("F2:F%d" % n, FormulaRule(formula=["F2>0"], fill=PatternFill("solid", bgColor="F4B6B6")))
        ws.conditional_formatting.add("G2:G%d" % n, FormulaRule(formula=["G2>0"], fill=PatternFill("solid", bgColor="FAD7A0")))
        ws.conditional_formatting.add("J2:J%d" % n, FormulaRule(formula=["J2>0"], fill=PatternFill("solid", bgColor="F4B6B6")))

    step_rows = []
    for s in report["sessions"]:
        prev_end = None
        for st in s["steps"]:
            if max_steps is not None and st["n"] > max_steps:
                break
            delta = (st["start"] - prev_end) if prev_end is not None else None
            prev_end = max(st["end"], prev_end or 0)
            step_rows.append([s["id"] or "(none)", st["n"], fmt_dt(st["start"], tz),
                              round(delta, 2) if delta is not None else None, st["tool"],
                              "; ".join(st["args"]), st["status"], status_text(st, lab),
                              round(st["dur"], 3) if st["dur"] is not None else None,
                              st["result"] or st["error"], st["reason"],
                              (st["submit"] or {}).get("result", "")])
    ws = _sheet(wb, lab["steps"], [lab["session"], "#", lab["time"], lab["delta"] + " (s)", lab["tool"],
                                   lab["args"], "status_code", lab["status"], lab["dur"] + " (s)",
                                   lab["result"], "reason_code", lab["submit"]],
                step_rows, [18, 6, 20, 12, 22, 60, 12, 22, 12, 60, 16, 10])
    if step_rows:
        n = len(step_rows) + 1
        rng = "A2:L%d" % n
        for code, color in (("error", "F4B6B6"), ("refused", "FAD7A0"), ("unavailable", "E1E5EA"),
                            ("no_outcome", "DCC8F2"), ("ok", "D5EFDD")):
            ws.conditional_formatting.add(rng, FormulaRule(formula=['$G2="%s"' % code],
                                                           fill=PatternFill("solid", bgColor=color)))
        ws.conditional_formatting.add("D2:D%d" % n, FormulaRule(
            formula=["AND(ISNUMBER(D2),D2>=%s)" % report["meta"]["gap_s"]], fill=PatternFill("solid", bgColor="F4B6B6")))

    gap_rows = []
    for s in report["sessions"]:
        for g in s["gaps"]:
            gap_rows.append([s["id"] or "(none)", fmt_dt(g["start"], tz), fmt_dt(g["end"], tz),
                             round(g["seconds"], 1), fmt_duration(g["shown_s"], lang), g["after_tool"],
                             g["before_tool"], g["other_sessions_calls"]])
    ws = _sheet(wb, lab["gaps"], [lab["session"], lab["first"], lab["last"], "seconds", lab["gap_line"],
                                  lab["gap_after"], lab["gap_before"], lab["other_calls"]],
                gap_rows, [18, 20, 20, 10, 14, 22, 22, 14])
    if gap_rows:
        ws.conditional_formatting.add("D2:D%d" % (len(gap_rows) + 1), ColorScaleRule(
            start_type="min", start_color="FDEBD0", end_type="max", end_color="E74C3C"))

    find_rows = [[fmt_dt(f["ts"], tz), f["session"] or "(none)", lab["kinds"].get(f["kind"], f["kind"]),
                  round(f["seconds"], 1) if f.get("seconds") else None, f.get("detail", "")]
                 for f in report["findings"]]
    _sheet(wb, lab["findings"], [lab["time"], lab["session"], lab["status"], "seconds", "detail"],
           find_rows, [20, 18, 36, 10, 60])

    tool_rows = [[t, v["calls"], v["ok"], v["error"], v["refused"], v["unavailable"], v["no_outcome"],
                  v["total_s"], v["max_s"]] for t, v in
                 sorted(report["stats"]["by_tool"].items(), key=lambda kv: -kv[1]["calls"])]
    tool_rows += [["", None, None, None, None, None, None, None, None]]
    tool_rows += [[lab["reasons"].get(k, k), v] for k, v in report["stats"]["reasons"].items()]
    ws = _sheet(wb, lab["by_tool"], [lab["tool"], lab["calls"], lab["ok"], lab["error"], lab["refused"],
                                     lab["unavailable"], lab["no_outcome"], "sum s", "max s"],
                tool_rows, [32, 9, 8, 8, 8, 10, 10, 10, 10])
    if tool_rows:
        ws.conditional_formatting.add("D2:D%d" % (len(tool_rows) + 1), FormulaRule(
            formula=["AND(ISNUMBER(D2),D2>0)"], fill=PatternFill("solid", bgColor="F4B6B6")))
    wb.save(path)
