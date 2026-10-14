"""Markdown and Excel reports built from aggregate.summarize() output."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Optional

from .score import ScoreRow


def _pct(x: Optional[float]) -> str:
    return "n/a" if x is None else "%.1f%%" % (100 * x)


def _num(x: Optional[float], nd: int = 3) -> str:
    return "n/a" if x is None else ("%." + str(nd) + "f") % x


def _sec(x: Optional[float]) -> str:
    return "n/a" if x is None else "%.1f" % x


def render_markdown(agg: dict, warnings: Iterable[str] = (), title: str = "Evaluation report",
                    notes: Iterable[str] = ()) -> str:
    out = ["# " + title, ""]
    warnings = list(warnings)
    if warnings:
        out += ["## Warnings", ""] + ["- " + w for w in warnings] + [""]
    for n in notes:
        out.append(n)
    if notes:
        out.append("")
    summary = agg["summary"]
    for split in agg["splits"]:
        label = "dev (tuning set: do not report as final)" if split == "dev" else "test (final evaluation)"
        out += ["## Split: " + label, ""]
        out.append("| System | Qs | Answered | Accuracy (95% CI) | Executed | Reproducibility | "
                   "Run variance | Confident wrong | Unhedged wrong | Interp. stated/asked/silent | "
                   "Median s |")
        out.append("|---|---|---|---|---|---|---|---|---|---|---|")
        for r in [s for s in summary if s["split"] == split]:
            ci = "n/a" if r["ci_lo"] is None else "%s [%s, %s]" % (
                _pct(r["accuracy"]), _pct(r["ci_lo"]), _pct(r["ci_hi"]))
            out.append("| %s | %d | %d | %s | %s | %s | %s | %d / %d | %d / %d | %d/%d/%d | %s |" % (
                r["system"], r["questions"], r["answered"], ci, _pct(r["executed_rate"]),
                _pct(r["reproducibility"]), _num(r["run_variance"], 4),
                r["confident_wrong"], r["wrong_total"], r["unhedged_wrong"], r["wrong_total"],
                r["interp_stated"], r["interp_asked"], r["interp_silent"], _sec(r["median_seconds"])))
        out.append("")
        errs = [s for s in summary if s["split"] == split]
        if errs:
            out.append("Error classes (failed answers): " + "; ".join(
                "%s: %s" % (s["system"], ", ".join("%s=%d" % (k, v) for k, v in s["errors"].items() if v) or "none")
                for s in errs))
            out.append("")
        diff = [q for q in agg["questions"] if q["split"] == split and q["differs"]]
        out.append("### Questions where systems differ (%d)" % len(diff))
        out.append("")
        if diff:
            out.append("| Question | " + " | ".join(agg["systems"]) + " |")
            out.append("|---|" + "---|" * len(agg["systems"]))
            for q in diff:
                cells = []
                for s in agg["systems"]:
                    v = q["systems"][s]
                    cells.append("not answered" if v["accuracy"] is None else
                                 "%s (%s)" % (_pct(v["accuracy"]), v.get("error_class", "none")))
                out.append("| %s | %s |" % (q["id"], " | ".join(cells)))
        else:
            out.append("None.")
        out.append("")
    out += [
        "## Definitions",
        "",
        "- Accuracy: share of gradable questions answered correctly (execution result matches the "
        "reference and all result checks pass), averaged over repeated runs per question. A question "
        "a system did not answer counts as 0. The interval is a percentile bootstrap over questions.",
        "- Reproducibility: average share of repeated runs that return the most common result "
        "(questions with at least two runs).",
        "- Run variance: sample variance of per-run accuracy across run numbers.",
        "- Confident wrong: executed, wrong, and the answer text states certainty. Unhedged wrong: "
        "executed, wrong, and no hedging language (includes answers that say nothing about confidence).",
        "",
    ]
    return "\n".join(out)


def write_excel(path: str | Path, agg: dict, rows: Iterable[ScoreRow],
                warnings: Iterable[str] = ()) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head = ["split", "system", "questions", "answered", "unanswered", "answers", "executed_rate",
            "accuracy", "ci_lo", "ci_hi", "run_variance", "runs", "reproducibility",
            "confident_wrong", "unhedged_wrong", "wrong_total", "interp_stated", "interp_asked",
            "interp_silent", "median_seconds"]
    ws.append(head)
    for r in agg["summary"]:
        ws.append([r[h] for h in head])
    _style_header(ws, Font, PatternFill, Alignment)
    n = ws.max_row
    if n >= 2:
        for col in ("G", "H", "I", "J", "M"):
            for rr in range(2, n + 1):
                ws["%s%d" % (col, rr)].number_format = "0.0%"
        ws.conditional_formatting.add(
            "H2:H%d" % n, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                          mid_type="num", mid_value=0.5, mid_color="FFEB84",
                                          end_type="num", end_value=1, end_color="63BE7B"))
        ws.conditional_formatting.add(
            "N2:O%d" % n, CellIsRule(operator="greaterThan", formula=["0"],
                                      fill=PatternFill("solid", start_color="F4B6B6", end_color="F4B6B6"),
                                      font=Font(bold=True, color="9C0006")))
    _autowidth(ws, get_column_letter)

    wq = wb.create_sheet("Questions")
    systems = agg["systems"]
    wq.append(["id", "split", "differs"] + ["%s accuracy" % s for s in systems] +
              ["%s reproducibility" % s for s in systems])
    for q in agg["questions"]:
        wq.append([q["id"], q["split"], "yes" if q["differs"] else "no"] +
                  [q["systems"][s]["accuracy"] for s in systems] +
                  [q["systems"][s]["reproducibility"] for s in systems])
    _style_header(wq, Font, PatternFill, Alignment)
    nq = wq.max_row
    if nq >= 2 and systems:
        last = get_column_letter(3 + len(systems))
        rng = "D2:%s%d" % (last, nq)
        for rr in range(2, nq + 1):
            for cc in range(4, 4 + 2 * len(systems)):
                wq.cell(row=rr, column=cc).number_format = "0%"
        wq.conditional_formatting.add(
            rng, ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                mid_type="num", mid_value=0.5, mid_color="FFEB84",
                                end_type="num", end_value=1, end_color="63BE7B"))
        wq.conditional_formatting.add(
            "C2:C%d" % nq, CellIsRule(operator="equal", formula=['"yes"'],
                                      fill=PatternFill("solid", start_color="FFE699", end_color="FFE699")))
    _autowidth(wq, get_column_letter)

    wa = wb.create_sheet("Answers")
    ah = ["id", "system", "run", "split", "executed", "match", "checks_pass", "correct", "error_class",
          "confidence", "confident_wrong", "unhedged_wrong", "interpretation", "seconds",
          "rows_returned", "detail", "error"]
    wa.append(ah)
    for r in rows:
        wa.append([getattr(r, h) if not isinstance(getattr(r, h), bool) else
                   ("yes" if getattr(r, h) else "no") for h in ah])
    _style_header(wa, Font, PatternFill, Alignment)
    na = wa.max_row
    if na >= 2:
        red = PatternFill("solid", start_color="F4B6B6", end_color="F4B6B6")
        green = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
        wa.conditional_formatting.add("H2:H%d" % na, CellIsRule(operator="equal", formula=['"no"'], fill=red))
        wa.conditional_formatting.add("H2:H%d" % na, CellIsRule(operator="equal", formula=['"yes"'], fill=green))
        wa.conditional_formatting.add("K2:K%d" % na, CellIsRule(operator="equal", formula=['"yes"'],
                                                                fill=red, font=Font(bold=True, color="9C0006")))
    _autowidth(wa, get_column_letter)

    warns = list(warnings)
    if warns:
        wn = wb.create_sheet("Warnings")
        wn.append(["warning"])
        for w in warns:
            wn.append([w])
        _style_header(wn, Font, PatternFill, Alignment)
        wn.column_dimensions["A"].width = 120
    wb.save(str(path))


def _style_header(ws, Font, PatternFill, Alignment) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="305496", end_color="305496")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _autowidth(ws, get_column_letter) -> None:
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
