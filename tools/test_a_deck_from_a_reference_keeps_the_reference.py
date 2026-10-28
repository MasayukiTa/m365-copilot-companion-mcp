# -*- coding: utf-8 -*-
"""A deck built from a reference must carry the reference, not start from blank.

WHY, in the operator's words: 「作成されたpptxがまったくもって過去のものを参照できていない」
「参照元があるならば参照元を直接中身の編集がなされるべき」「idでもてるは初期構想としてmust」.
create_pptx opens an empty python-pptx template, so its output inherits nothing from the
material it was supposedly based on. pptx_from_reference instead OPENS the source deck and
rewrites only the named text, so the source's design (here stood in for by a distinctive theme
colour and an extra slide the blank template never has) survives into the product. These tests
pin exactly that: the design is carried, the text is replaced, the source id resolves a deck,
and the product points back at its source.
"""
from __future__ import annotations

import os
import sys

import pytest

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO not in sys.path:
    sys.path.insert(0, REPO)

pytest.importorskip("pptx", reason="python-pptx is what opens and preserves the reference deck")

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402

from tools import pptx_ops, pptx_provenance as prov  # noqa: E402

# A theme colour the blank python-pptx template never uses, so finding it in the output proves
# the output came from the source deck rather than a fresh Presentation().
MARKER_RGB = RGBColor(0xAB, 0x12, 0xEF)


@pytest.fixture(autouse=True)
def _unlocked(monkeypatch):
    """Pass the mutating-tool gate in-process.

    require_unlocked() denies in-process callers (no HTTP context) unless the owner's switch is
    off. conftest deliberately defaults the switch ON so gate tests see the real behaviour, so a
    test that drives a mutating tool must turn it off for itself -- the switch's own docstring
    names this exact exit. Patch the switch, not a password: unlock() cannot run without the
    same HTTP context this branch is missing.
    """
    from tools import unlock_switch
    monkeypatch.setattr(unlock_switch, "unlock_required", lambda: False)


def _unlock():
    """No-op: the autouse _unlocked fixture already opened the gate for every test."""
    return None


@pytest.fixture
def reference_deck(tmp_path):
    """A source deck with a distinctive theme colour and two slides of real text."""
    p = tmp_path / "past_report.pptx"
    prs = Presentation()
    # Slide 1: title + body, title carries the marker colour.
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "AUGUST REPORT"
    run = s1.shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = MARKER_RGB
    for ph in s1.placeholders:
        if ph != s1.shapes.title and ph.has_text_frame:
            ph.text_frame.text = "old revenue line"
            break
    # Slide 2: an extra slide a blank template would never have.
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "AUGUST APPENDIX"
    prs.save(str(p))
    return p


def test_the_output_keeps_the_source_decks_design_not_a_blank_template(reference_deck, tmp_path):
    _unlock()
    out = tmp_path / "september.pptx"
    msg = pptx_ops.pptx_from_reference(str(out), str(reference_deck))
    assert out.is_file(), msg
    prs = Presentation(str(out))
    # The extra slide only the source had is present -> not built from an empty template.
    assert len(prs.slides) == 2, msg
    # The source's marker theme colour survived into the product.
    run = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == MARKER_RGB


def test_named_text_is_replaced_in_place(reference_deck, tmp_path):
    _unlock()
    out = tmp_path / "september.pptx"
    updates = [
        {"slide": 1, "title": "SEPTEMBER REPORT", "replace": {"old revenue line": "new revenue line"}},
    ]
    msg = pptx_ops.pptx_from_reference(str(out), str(reference_deck), updates=updates)
    assert out.is_file(), msg
    prs = Presentation(str(out))
    assert prs.slides[0].shapes.title.text == "SEPTEMBER REPORT"
    body_texts = [
        ph.text_frame.text
        for ph in prs.slides[0].placeholders
        if ph != prs.slides[0].shapes.title and ph.has_text_frame
    ]
    assert any("new revenue line" in t for t in body_texts)
    assert all("old revenue line" not in t for t in body_texts)
    # The source's extra slide and its title survived unchanged -- only slide 1 was edited.
    assert len(prs.slides) == 2
    assert prs.slides[1].shapes.title.text == "AUGUST APPENDIX"


def test_the_source_file_on_disk_is_not_modified(reference_deck, tmp_path):
    _unlock()
    before = reference_deck.read_bytes()
    out = tmp_path / "september.pptx"
    pptx_ops.pptx_from_reference(
        str(out), str(reference_deck), updates=[{"slide": 1, "title": "CHANGED"}]
    )
    assert reference_deck.read_bytes() == before


def test_a_reference_given_as_a_provenance_id_resolves_the_right_deck(reference_deck, tmp_path):
    _unlock()
    # Stamp the source with a run id, then address it by that id alone (no path).
    run_id = "r_test_refedit_0001"
    prov.stamp(str(reference_deck), sources=[], run_id=run_id)
    out = tmp_path / "september.pptx"
    msg = pptx_ops.pptx_from_reference(
        str(out), run_id, updates=[{"slide": 1, "title": "BY ID"}], search_dir=str(tmp_path)
    )
    assert out.is_file(), msg
    prs = Presentation(str(out))
    assert prs.slides[0].shapes.title.text == "BY ID"
    assert len(prs.slides) == 2


def test_an_unknown_id_is_an_explicit_error_not_a_blank_deck(reference_deck, tmp_path):
    _unlock()
    out = tmp_path / "september.pptx"
    msg = pptx_ops.pptx_from_reference(
        str(out), "r_no_such_run", updates=[{"slide": 1, "title": "X"}], search_dir=str(tmp_path)
    )
    assert not out.is_file()
    assert "error" in msg.lower() and "provenance id" in msg.lower()


def test_the_product_points_back_at_its_source(reference_deck, tmp_path):
    _unlock()
    out = tmp_path / "september.pptx"
    pptx_ops.pptx_from_reference(
        str(out), str(reference_deck), updates=[{"slide": 1, "title": "TRACED"}],
        run_id="r_prod_0007",
    )
    tags = prov.read(str(out))
    # Agent-made, this run, and CompanionSources carries the source's digest.
    assert tags.get(prov.MADE) == "1"
    assert tags.get(prov.RUN_ID) == "r_prod_0007"
    src_digest = prov.digest(str(reference_deck))
    assert src_digest and src_digest in tags.get(prov.SOURCES, "")
