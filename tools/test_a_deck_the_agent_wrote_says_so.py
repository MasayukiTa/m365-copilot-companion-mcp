# -*- coding: utf-8 -*-
"""A deck the Copilot agent wrote must say the agent wrote it.

MEASURED 2026-09-15 on a deck a worker produced that morning from a September 2nd template:
`python-pptx` carries the template's core properties across, so the file reported the operator
as its author, the operator as its last editor, and a creation time two weeks before it existed.
It did not merely fail to say a machine made it -- it said a person did.

WHAT THAT COST. A worker asked for "this month's report from the past material" finds last
week's agent-made report in the same folder and treats it as past material. That happened: the
first of four OGF runs reported the work complete by pointing at output an earlier run had made,
and the holdout that made the measurement mean anything had to be assembled by hand, moving 94
files aside to decide which were source and which were product.

BINARY, NOT A GENERATION COUNT. An earlier draft of this recorded how many machine generations a
deck stood from a human original. That line does not exist in this folder: the operator's own
August and September material was drafted by GPT or Claude and corrected by hand. What the
folder does contain is Copilot-agent output, and telling that apart from everything else is the
distinction worth keeping.

THE STAMP IS NOT A BOUNDARY, and the tests below hold that too. A deck written with `zipfile` by
hand, or by a different interpreter, carries nothing. `made_by_agent` is a positive signal only.
"""
from __future__ import annotations

import io
import os
import subprocess
import sys
import zipfile

import pytest

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO not in sys.path:
    sys.path.insert(0, REPO)

from tools import _subproc, pptx_provenance as prov  # noqa: E402

pytest.importorskip("pptx", reason="python-pptx is what carries the stale properties")


@pytest.fixture
def deck(tmp_path):
    """A minimal real .pptx, made the way a worker makes one."""
    from pptx import Presentation
    p = tmp_path / "source.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(p))
    return p


def test_an_unstamped_deck_does_not_claim_the_agent_made_it(deck):
    assert prov.read(str(deck)) == {}
    assert prov.made_by_agent(str(deck)) is False


def test_stamping_makes_it_say_so_and_leaves_the_deck_openable(deck):
    from pptx import Presentation
    before = len(Presentation(str(deck)).slides)
    prov.stamp(str(deck), sources=[], run_id="r123")
    assert prov.made_by_agent(str(deck)) is True
    assert prov.read(str(deck))[prov.RUN_ID] == "r123"
    assert len(Presentation(str(deck)).slides) == before, "the tag changed the deck"
    with zipfile.ZipFile(str(deck)) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert prov.CUSTOM_PART in names
        ct = z.read("[Content_Types].xml").decode("utf-8", "replace")
        assert "/docProps/custom.xml" in ct, (
            "an undeclared part is one PowerPoint refuses to open the file over")
        rels = z.read("_rels/.rels").decode("utf-8", "replace")
        assert "docProps/custom.xml" in rels, "an orphaned part is one nothing will read"


def test_a_property_the_operator_set_is_not_destroyed(deck):
    """The part is shared. PowerPoint writes custom properties there too."""
    prov.stamp(str(deck), sources=[], run_id="r1")
    # Add a foreign property the way another producer would, then stamp again.
    with zipfile.ZipFile(str(deck)) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    merged = prov._parse_custom(
        [d for i, d in items if i.filename == prov.CUSTOM_PART][0].decode("utf-8"))
    merged["ContentTypeId"] = "0x0101"
    payload = prov._render_custom(merged)
    tmp = str(deck) + ".t"
    with zipfile.ZipFile(tmp, "w") as out:
        for i, d in items:
            out.writestr(i, payload if i.filename == prov.CUSTOM_PART else d)
    os.replace(tmp, str(deck))

    prov.stamp(str(deck), sources=[], run_id="r2")
    with zipfile.ZipFile(str(deck)) as z:
        props = prov._parse_custom(z.read(prov.CUSTOM_PART).decode("utf-8"))
    assert props.get("ContentTypeId") == "0x0101", "another producer's property was dropped"
    assert props[prov.RUN_ID] == "r2"


def test_the_sources_it_was_built_from_are_recorded(deck, tmp_path):
    other = tmp_path / "base.pptx"
    io.open(str(other), "wb").write(io.open(str(deck), "rb").read())
    prov.stamp(str(deck), sources=[str(other)])
    assert "base.pptx@" in prov.read(str(deck))[prov.SOURCES]


def test_a_failed_stamp_leaves_the_original_where_it_was(tmp_path):
    """The tag must never cost the deck. A truncated archive mistaken for a backup is how this
    repository lost 9.5GB of snapshots on 2026-09-14."""
    broken = tmp_path / "not_really.pptx"
    broken.write_bytes(b"not a zip at all")
    with pytest.raises(Exception):
        prov.stamp(str(broken))
    assert broken.read_bytes() == b"not a zip at all"


# ── the choke point ────────────────────────────────────────────────────────────────────────

def test_the_child_environment_carries_the_stamp_module():
    env = _subproc.sanitized_child_env()
    first = (env.get("PYTHONPATH") or "").split(os.pathsep)[0]
    assert first.endswith("_pptx_autostamp"), (
        "worker code runs in a child interpreter; nothing else reaches code it composed itself")
    assert env.get("MCP_COMPANION_REPO"), "the child has no other way to find the stamp module"


def test_the_kill_switch_is_read_from_the_environment_being_built():
    """Checking os.environ alone looked equivalent and was not: a caller handing in its own
    environment -- which is how this is tested -- would have the switch ignored."""
    off = _subproc._with_pptx_autostamp({"MCP_NO_PPTX_AUTOSTAMP": "1"})
    assert "_pptx_autostamp" not in (off.get("PYTHONPATH") or "")


def test_worker_style_code_that_knows_nothing_about_stamping_still_stamps(deck, tmp_path):
    """END TO END, in the shape that actually occurs. Across four OGF runs, every deck was
    written by python-pptx code the worker composed and handed to run_python -- not one used
    the fleet's own deck-writing tools, so a stamp in those would have covered none of it."""
    out = tmp_path / "written_by_worker.pptx"
    script = tmp_path / "w.py"
    script.write_text(
        "from pptx import Presentation\n"
        "prs = Presentation(r'%s')\n"
        "prs.save(r'%s')\n" % (str(deck), str(out)), encoding="utf-8")

    env = _subproc.sanitized_child_env()
    env["MCP_FLEET_RUN_ID"] = "rZZZ"
    r = subprocess.run([sys.executable, str(script)], env=env, capture_output=True, timeout=180)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")[:400]

    assert prov.made_by_agent(str(out)) is True, "worker-written deck carries no tag"
    tag = prov.read(str(out))
    assert tag[prov.RUN_ID] == "rZZZ"
    assert "source.pptx@" in tag[prov.SOURCES], "the deck it was built from was not recorded"
    from pptx import Presentation
    assert Presentation(str(out)) is not None


def test_a_deck_is_not_recorded_as_its_own_source(deck, tmp_path):
    """Opening a deck and saving over it is the commonest edit there is."""
    script = tmp_path / "w.py"
    script.write_text(
        "from pptx import Presentation\n"
        "prs = Presentation(r'%s')\n"
        "prs.save(r'%s')\n" % (str(deck), str(deck)), encoding="utf-8")
    env = _subproc.sanitized_child_env()
    r = subprocess.run([sys.executable, str(script)], env=env, capture_output=True, timeout=180)
    assert r.returncode == 0, r.stderr.decode("utf-8", "replace")[:400]
    assert "source.pptx@" not in prov.read(str(deck)).get(prov.SOURCES, "")
