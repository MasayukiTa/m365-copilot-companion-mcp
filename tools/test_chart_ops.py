"""chart_render / excel_chart_export / chart_from_python (tools/chart_ops.py).

Generic data-to-chart tools. These tests run on Linux CI: no Windows paths, no real unlock
state (the gate is exercised through the same HTTP-context patch the other gate tests use).
"""
from __future__ import annotations

import sys
import textwrap
from unittest.mock import MagicMock, patch

import pytest

pytest.importorskip("matplotlib")
pytest.importorskip("openpyxl")
pytest.importorskip("xlsxwriter")
pytest.importorskip("pandas")

import openpyxl  # noqa: E402

from tools import chart_ops as C  # noqa: E402
from tools import file_ops  # noqa: E402

PNG = b"\x89PNG\r\n\x1a\n"

ROWS = [
    {"部門": "営業", "月": "1", "売上": 120, "原価": 70},
    {"部門": "営業", "月": "2", "売上": 150, "原価": 80},
    {"部門": "開発", "月": "1", "売上": 90, "原価": 60},
    {"部門": "開発", "月": "2", "売上": 110, "原価": 65},
    {"部門": "総務", "月": "1", "売上": 40, "原価": 30},
    {"部門": "総務", "月": "2", "売上": 45, "原価": 33},
]


@pytest.fixture(autouse=True)
def _open_gate_and_base(tmp_path, monkeypatch):
    """Unlocked, and tmp_path is an allowed base (the gate itself is tested separately)."""
    monkeypatch.setattr(C, "require_unlocked", lambda: None)
    monkeypatch.setattr(file_ops, "ALLOWED_BASES", [tmp_path.resolve()])
    try:
        from tools import folder_policy
        monkeypatch.setattr(folder_policy, "allowed_bases", lambda: None)
    except Exception:
        pass
    yield


def _is_png(path) -> bool:
    with open(path, "rb") as f:
        return f.read(8) == PNG


# ------------------------------------------------------------------ chart_render


@pytest.mark.parametrize("kind", list(C.CHART_KINDS))
def test_every_kind_makes_a_png(tmp_path, kind):
    out = tmp_path / f"{kind}.png"
    x = "売上" if kind == "scatter" else "部門"
    y = "原価" if kind == "scatter" else "売上"
    res = C.chart_render(kind, rows=ROWS, x=x, y=y, title="部門別の売上", out_path=str(out))
    assert not res.startswith("["), res
    assert out.stat().st_size > 0 and _is_png(out)
    assert f"kind={kind}" in res and "points=" in res


def test_japanese_title_and_labels_do_not_crash(tmp_path):
    out = tmp_path / "jp.png"
    res = C.chart_render("bar", rows=ROWS, x="部門", y=["売上", "原価"], title="日本語タイトル",
                         xlabel="部門", ylabel="金額", out_path=str(out))
    assert not res.startswith("["), res
    assert _is_png(out)
    assert "売上, 原価" in res


def test_group_makes_one_series_per_value(tmp_path):
    res = C.chart_render("line", rows=ROWS, x="月", y="売上", group="部門",
                         out_path=str(tmp_path / "g.png"))
    assert "series (3)" in res
    for name in ("営業", "開発", "総務"):
        assert name in res
    res = C.chart_render("stacked_bar", rows=ROWS, x="月", y="売上", group="部門",
                         out_path=str(tmp_path / "s.png"))
    assert "series (3)" in res


def test_top_n_and_sort(tmp_path):
    res = C.chart_render("bar", rows=ROWS, x="部門", y="売上", top_n=2, sort="desc",
                         out_path=str(tmp_path / "t.png"))
    assert "top_n=2: kept 2 of 3" in res


def test_row_limit_is_applied_and_stated(tmp_path):
    csv = tmp_path / "big.csv"
    csv.write_text("k,v\n" + "\n".join(f"a{i},{i}" for i in range(500)), encoding="utf-8")
    res = C.chart_render("line", data_path=str(csv), x="v", y="v", max_rows=50,
                         out_path=str(tmp_path / "big.png"))
    assert "truncated to the first 50 rows" in res
    assert "input rows used: 50" in res


def test_csv_and_xlsx_input_with_sheet_and_range(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "メモ"
    ws.append(["ignore"])
    ws2 = wb.create_sheet("data")
    ws2.append(["junk", "junk"])
    ws2.append(["name", "val"])
    ws2.append(["a", 1])
    ws2.append(["b", 2])
    ws2.append(["c", 3])
    p = tmp_path / "in.xlsx"
    wb.save(p)
    res = C.chart_render("bar", data_path=str(p), sheet="data", cell_range="A2:B5",
                         x="name", y="val", out_path=str(tmp_path / "x.png"))
    assert not res.startswith("["), res
    assert "points=3" in res


def test_bad_inputs_are_errors_not_crashes(tmp_path):
    assert C.chart_render("nope", rows=ROWS).startswith("[chart_render error")
    assert "not found" in C.chart_render("bar", rows=ROWS, x="無い列", y="売上")
    assert C.chart_render("bar", rows=ROWS, x="部門", y="売上",
                          out_path=str(tmp_path / "a.txt")).startswith("[chart_render error")
    assert C.chart_render("bar").startswith("[chart_render error")


def test_missing_dependency_is_a_setup_message(monkeypatch):
    def boom(name):
        raise ImportError("No module named " + name)
    monkeypatch.setattr("importlib.import_module", boom)
    res = C.chart_render("bar", rows=ROWS, x="部門", y="売上")
    assert res.startswith("[chart_render error")
    assert "requirements.txt の依存が未導入です" in res and "setup/quickstart" in res


# ------------------------------------------------------------------ excel_chart_export


def _chart_refs(chart) -> list[str]:
    refs = []
    for s in chart.series:
        if s.val is not None and s.val.numRef is not None:
            refs.append(s.val.numRef.f)
        if s.cat is not None:
            ref = s.cat.numRef or s.cat.strRef
            if ref is not None:
                refs.append(ref.f)
    return refs


def test_excel_roundtrip_data_chart_and_notes(tmp_path):
    out = tmp_path / "r.xlsx"
    res = C.excel_chart_export("column", rows=ROWS, x="部門", y=["売上", "原価"],
                               title="部門別", out_path=str(out), notes="メモです",
                               request="部門別の売上を出して", source="inline sample")
    assert not res.startswith("["), res
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[:3] == ["データ", "グラフ", "説明"]
    data = wb["データ"]
    header = [c.value for c in data[1]]
    assert header == ["部門", "月", "売上", "原価"]
    for i, row in enumerate(ROWS, start=2):
        got = [c.value for c in data[i]]
        assert got[0] == row["部門"] and got[2] == row["売上"] and got[3] == row["原価"]
    assert data.auto_filter.ref  # autofilter on the table
    assert data.max_row == len(ROWS) + 1

    # a native chart object, whose series refer to data-sheet cells (not a pasted picture)
    charts = wb["グラフ"]._charts
    assert len(charts) == 1
    refs = _chart_refs(charts[0])
    assert refs and all("データ" in r for r in refs), refs
    assert any("$C$2:$C$7" in r for r in refs), refs   # 売上 column
    assert any("$A$2:$A$7" in r for r in refs), refs   # 部門 categories
    assert not wb["グラフ"]._images

    texts = [c.value for row in wb["説明"].iter_rows() for c in row if c.value]
    assert "メモです" in texts and "部門別の売上を出して" in texts and "inline sample" in texts
    assert any("6 行" in str(t) for t in texts)


@pytest.mark.parametrize("kind", list(C.EXCEL_KINDS) + ["stacked"])
def test_every_excel_kind_has_a_chart(tmp_path, kind):
    out = tmp_path / f"{kind}.xlsx"
    res = C.excel_chart_export(kind, rows=ROWS, x="部門", y="売上", out_path=str(out))
    assert not res.startswith("["), res
    wb = openpyxl.load_workbook(out)
    assert len(wb["グラフ"]._charts) == 1


def test_group_goes_through_a_summary_sheet_but_still_references_cells(tmp_path):
    out = tmp_path / "g.xlsx"
    res = C.excel_chart_export("stacked", rows=ROWS, x="月", y="売上", group="部門",
                               out_path=str(out))
    assert not res.startswith("["), res
    wb = openpyxl.load_workbook(out)
    assert "集計" in wb.sheetnames and "データ" in wb.sheetnames
    refs = _chart_refs(wb["グラフ"]._charts[0])
    assert refs and all("集計" in r for r in refs)
    assert len(wb["グラフ"]._charts[0].series) == 3


def test_existing_file_is_not_overwritten(tmp_path):
    out = tmp_path / "keep.xlsx"
    out.write_bytes(b"precious")
    res = C.excel_chart_export("line", rows=ROWS, x="月", y="売上", out_path=str(out))
    assert out.read_bytes() == b"precious"
    second = tmp_path / "keep_2.xlsx"
    assert second.exists() and str(second) in res
    res = C.excel_chart_export("line", rows=ROWS, x="月", y="売上", out_path=str(out))
    assert (tmp_path / "keep_3.xlsx").exists() and out.read_bytes() == b"precious"
    res = C.excel_chart_export("line", rows=ROWS, x="月", y="売上", out_path=str(out),
                               overwrite=True)
    assert not res.startswith("[") and out.read_bytes() != b"precious"
    assert not list(tmp_path.glob("*.tmp"))


def test_excel_bad_inputs(tmp_path):
    assert C.excel_chart_export("hist", rows=ROWS).startswith("[excel_chart_export error")
    assert C.excel_chart_export("column", rows=ROWS, out_path=str(tmp_path / "a.csv")
                                ).startswith("[excel_chart_export error")


# ------------------------------------------------------------------ chart_from_python

GOOD_CODE = textwrap.dedent("""
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots()
    ax.bar(["あ", "い"], [1, 2])
    ax.set_title("日本語")
    fig.savefig(OUT_PATH)
    print("drew")
""")


def test_python_chart_makes_a_png(tmp_path):
    out = tmp_path / "py.png"
    res = C.chart_from_python(GOOD_CODE, out_path=str(out))
    assert not res.startswith("["), res
    assert _is_png(out) and "drew" in res


def test_python_chart_autosaves_an_open_figure_and_sees_data_path(tmp_path):
    csv = tmp_path / "d.csv"
    csv.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    code = ("import pandas as pd, matplotlib.pyplot as plt\n"
            "df = pd.read_csv(DATA_PATH)\nplt.plot(df['a'], df['b'])\n")
    out = tmp_path / "auto.png"
    res = C.chart_from_python(code, data_path=str(csv), out_path=str(out))
    assert not res.startswith("["), res
    assert _is_png(out)


def test_python_chart_timeout_kills_the_child(tmp_path):
    out = tmp_path / "slow.png"
    res = C.chart_from_python("import time\ntime.sleep(60)\n", out_path=str(out), timeout=2)
    assert res.startswith("[chart_from_python error") and "timed out" in res
    assert not out.exists()


def test_python_chart_without_png_is_an_error(tmp_path):
    res = C.chart_from_python("print('no figure')\n", out_path=str(tmp_path / "n.png"))
    assert res.startswith("[chart_from_python error") and "no PNG" in res
    assert "no figure" in res


def test_python_chart_exception_returns_stderr_tail(tmp_path):
    res = C.chart_from_python("raise ValueError('kaboom')\n", out_path=str(tmp_path / "e.png"))
    assert res.startswith("[chart_from_python error") and "kaboom" in res


def test_python_chart_does_not_overwrite(tmp_path):
    out = tmp_path / "o.png"
    out.write_bytes(b"precious")
    res = C.chart_from_python(GOOD_CODE, out_path=str(out))
    assert out.read_bytes() == b"precious" and (tmp_path / "o_2.png").exists()
    assert "o_2.png" in res


def test_python_chart_workdir_is_cleaned_up(tmp_path):
    from relay.temp_home import temp_dir
    root = temp_dir("chart_from_python")
    import os
    before = set(os.listdir(root))
    C.chart_from_python(GOOD_CODE, out_path=str(tmp_path / "c.png"))
    assert set(os.listdir(root)) == before


# ------------------------------------------------------------------ the gate


def _locked_request() -> MagicMock:
    req = MagicMock()
    req.client = MagicMock()
    req.client.host = "198.51.100.77"
    req.headers = MagicMock()
    req.headers.get = lambda key, default="": default
    return req


@pytest.mark.parametrize("call", [
    lambda: C.chart_render("bar", rows=ROWS, x="部門", y="売上"),
    lambda: C.excel_chart_export("column", rows=ROWS, x="部門", y="売上"),
    lambda: C.chart_from_python("print(1)"),
])
def test_all_three_refuse_when_locked(tmp_path, monkeypatch, call):
    from tools import security as sec
    monkeypatch.setattr(sec, "STATE_FILE", tmp_path / ".unlock_state.json")
    # the module-level patch from the autouse fixture is undone: use the real gate
    from tools.security import require_unlocked as real
    monkeypatch.setattr(C, "require_unlocked", real)
    with patch("tools.security.get_http_request", return_value=_locked_request()):
        result = call()
    assert result.startswith("[locked")
    assert not list(tmp_path.glob("*.png")) and not list(tmp_path.glob("*.xlsx"))


def test_the_three_tools_are_in_the_gateway_catalogue():
    from tools import tool_catalogue
    assert tool_catalogue.categorise("chart_render") == "data"
    assert tool_catalogue.categorise("excel_chart_export") == "data"
    assert tool_catalogue.categorise("chart_from_python") == "run"
    assert sys.version_info >= (3, 10)


# ------------------------------------------------------------------ chart_render: interval

INTERVAL_ROWS = [
    {"item": "item A", "est": 0.8, "lo": 0.5, "hi": 1.1, "grp": "g1"},
    {"item": "item B", "est": 0.1, "lo": -0.2, "hi": 0.4, "grp": "g2"},
    {"item": "item C", "est": 0.3, "lo": None, "hi": None, "grp": "g2"},
]


def _interval(tmp_path, name="iv.png", **kw):
    args = dict(rows=INTERVAL_ROWS, x="est", y="item", lo="lo", hi="hi", group="grp",
                out_path=str(tmp_path / name))
    args.update(kw)
    return C.chart_render("interval", **args)


def test_interval_makes_a_png_with_groups_and_a_missing_interval(tmp_path):
    res = _interval(tmp_path, ref="0,0.3,-0.3", palette="g1=#0f8b8d,g2=#e8871e",
                    hollow_groups="g2", title="区間つきの点")
    assert not res.startswith("["), res
    assert _is_png(tmp_path / "iv.png")
    assert "kind=interval" in res and "points=3" in res and "series (2)" in res
    assert "1 row(s) without an interval, marked '区間なし'" in res


def test_interval_custom_na_label_is_reported(tmp_path):
    res = _interval(tmp_path, na_label="n/a")
    assert "marked 'n/a'" in res


def test_interval_without_group_or_interval_columns(tmp_path):
    res = _interval(tmp_path, group="", lo="", hi="")
    assert not res.startswith("["), res
    assert "points=3" in res and "without an interval" not in res
    assert "series (1)" in res


def test_interval_without_group_still_draws_bars(tmp_path):
    res = _interval(tmp_path, group="")
    assert not res.startswith("["), res
    assert "1 row(s) without an interval" in res


@pytest.mark.parametrize("palette", ["g1", "g1=red", "=#112233", "g1=#12345", "g1=#0f8b8d,g2"])
def test_interval_invalid_palette_is_an_error(tmp_path, palette):
    res = _interval(tmp_path, palette=palette)
    assert res.startswith("[chart_render error:") and "palette" in res
    assert not (tmp_path / "iv.png").exists()


@pytest.mark.parametrize("ref", ["", " ", ",", None])
def test_interval_empty_ref_draws_no_reference_lines(tmp_path, ref):
    assert C._parse_ref(ref) == []
    res = _interval(tmp_path, ref=ref)
    assert not res.startswith("["), res


def test_interval_bad_ref_is_an_error(tmp_path):
    assert _interval(tmp_path, ref="0,abc").startswith("[chart_render error:")
    assert _interval(tmp_path, ref="nan").startswith("[chart_render error:")


def test_parse_palette_reads_pairs():
    assert C._parse_palette("g1=#0f8b8d, g2=#E8871E") == {"g1": "#0f8b8d", "g2": "#E8871E"}
    assert C._parse_palette("") == {}


def test_interval_hollow_groups_draw_white_filled_points(tmp_path):
    from matplotlib.figure import Figure
    pd = pytest.importorskip("pandas")
    fig = Figure()
    ax = fig.add_subplot(111)
    C._draw_interval(ax, pd, pd.DataFrame(INTERVAL_ROWS), "est", "item", "lo", "hi", "grp",
                     [], "none", {"g1": "#0f8b8d", "g2": "#e8871e"}, {"g2"})
    faces = {ln.get_markeredgecolor(): ln.get_markerfacecolor() for ln in ax.lines
             if ln.get_marker() == "o"}
    assert faces["#0f8b8d"] == "#0f8b8d"
    assert faces["#e8871e"] == "white"


def test_interval_y_axis_has_margin_and_first_row_on_top():
    from matplotlib.figure import Figure
    pd = pytest.importorskip("pandas")
    fig = Figure()
    ax = fig.add_subplot(111)
    C._draw_interval(ax, pd, pd.DataFrame(INTERVAL_ROWS), "est", "item", "lo", "hi", "grp",
                     [0.0], "none", {}, set())
    top, bottom = ax.get_ylim()[1], ax.get_ylim()[0]
    n = len(INTERVAL_ROWS)
    assert top <= -0.5 and bottom >= n - 1 + 0.5      # rows 0 and n-1 stay off the axes
    assert top < bottom                                   # inverted: row 0 is drawn first (top)
    assert [t.get_text() for t in ax.get_yticklabels()] == ["item A", "item B", "item C"]


def test_interval_argument_errors(tmp_path):
    assert "lo and hi together" in _interval(tmp_path, hi="")
    assert "needs x" in _interval(tmp_path, x="")
    assert "column 'nope' not found" in _interval(tmp_path, lo="nope")


def test_existing_kinds_ignore_the_interval_arguments(tmp_path):
    a = C.chart_render("bar", rows=ROWS, x="部門", y="売上", out_path=str(tmp_path / "a.png"))
    b = C.chart_render("bar", rows=ROWS, x="部門", y="売上", out_path=str(tmp_path / "b.png"),
                       lo="", hi="", ref="", palette="", hollow_groups="")
    assert not a.startswith("[") and not b.startswith("[")
    assert a.splitlines()[1:] == b.splitlines()[1:]
    assert (tmp_path / "a.png").read_bytes() == (tmp_path / "b.png").read_bytes()
    assert C.CHART_KINDS[:7] == ("bar", "line", "scatter", "hist", "pie", "stacked_bar", "box")
