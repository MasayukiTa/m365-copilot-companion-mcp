"""Reader feedback: the OK/NG column of the summary sheet -> feedback.jsonl (appended).

feedback.jsonl lives in the PARENT of the run folder, so it collects every night in one place.
Each line carries the item id and the SQL shown in the report, so a later step (a library of
confirmed queries) can pick up the verified ones. No such step is implemented here.
"""
from __future__ import annotations

import json
from pathlib import Path

from tools.overnight.common import Clock, fmt_dt

OK_MARKS = {"○", "〇", "◯", "o", "ok", "✓", "✔", "good", "yes", "y"}
NG_MARKS = {"×", "✗", "✕", "x", "ng", "no", "n", "bad"}
FILENAME = "feedback.jsonl"


def verdict_of(mark) -> str:
    if mark is None:
        return ""
    m = str(mark).strip().lower()
    if m in OK_MARKS:
        return "ok"
    if m in NG_MARKS:
        return "ng"
    return ""


def read_marks(report: Path) -> list:
    from openpyxl import load_workbook
    wb = load_workbook(report, read_only=True, data_only=True)
    try:
        ws = wb["サマリ"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    h = next((i for i, row in enumerate(rows[:3]) if "OK/NG" in [str(x or "").strip() for x in row]), 0)
    head = [str(x or "").strip() for x in rows[h]]
    col = {name: i for i, name in enumerate(head)}
    need = ("ID", "OK/NG")
    if any(n not in col for n in need):
        return []
    out = []
    for r in rows[h + 1:]:
        iid = r[col["ID"]]
        v = verdict_of(r[col["OK/NG"]])
        if not iid or not v:
            continue
        get = lambda name: (r[col[name]] if name in col and col[name] < len(r) else "") or ""
        out.append({"id": str(iid), "ask": str(get("依頼文")), "result": str(get("結果")),
                    "verdict": v, "sql": str(get("SQL(参考)"))})
    return out


def _existing(path: Path) -> set:
    seen = set()
    try:
        for ln in path.read_text(encoding="utf-8").splitlines():
            try:
                d = json.loads(ln)
                seen.add((d.get("report"), d.get("id"), d.get("verdict")))
            except Exception:
                continue
    except OSError:
        pass
    return seen


def ingest(report, clock=None) -> int:
    """Append new marks from one report; returns how many lines were added."""
    report = Path(report)
    target = report.parent.parent / FILENAME
    seen = _existing(target)
    marks = read_marks(report)
    now = fmt_dt((clock or Clock()).now())
    added = 0
    with open(target, "a", encoding="utf-8") as fh:
        for m in marks:
            key = (str(report), m["id"], m["verdict"])
            if key in seen:
                continue
            m.update({"ts": now, "report": str(report)})
            fh.write(json.dumps(m, ensure_ascii=False) + "\n")
            added += 1
    return added


def ingest_siblings(out_dir, log=None) -> int:
    """At the start of a run: pick up marks from earlier nights' reports. Never raises."""
    total = 0
    try:
        for rep in sorted(Path(out_dir).parent.glob("*/朝レポート_*.xlsx")):
            if rep.parent == Path(out_dir):
                continue
            try:
                total += ingest(rep)
            except Exception as e:
                if log:
                    log("feedback: could not read %s: %r" % (rep, e))
    except Exception:
        pass
    return total
