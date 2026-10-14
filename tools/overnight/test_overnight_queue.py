# -*- coding: utf-8 -*-
"""Queue route, 'could not be submitted' reporting, and the morning consistency check."""
from __future__ import annotations

import json
import os
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from tools.overnight import batch_cli as cli
from tools.overnight import classify, reconcile, report
from tools.overnight import state as S
from tools.overnight.common import JST
from tools.overnight.detect import DONE, FAILED, QUEUED, RUNNING, UNSEEN, FleetReader, Observation
from tools.overnight.runner import Runner
from tools.overnight.submit import QueueSubmitter, SubmitResult


# ------------------------------------------------------------------ QueueSubmitter
def test_queue_submitter_reads_the_job_id_from_the_door_reply():
    q = QueueSubmitter(fn=lambda g, note="", source="": "queued abc123def456 -- the goal is in this machine's queue. 1 job(s) waiting.")
    r = q.submit("[A1] 件数を数えて。保存して。")
    assert r.ok and r.job_id == "abc123def456" and not r.duplicate


def test_queue_submitter_maps_the_similar_goal_refusal_to_duplicate():
    msg = ("[fleet_submit: refused -- this reads as the goal already queued or running "
           "as abc123 (70% word overlap). If it is genuinely different work, say specifically]")
    r = QueueSubmitter(fn=lambda *a, **k: msg).submit("[A1] x。")
    assert r.duplicate and not r.ok


@pytest.mark.parametrize("msg", [
    "[fleet_submit: refused -- 50 jobs are already waiting and nothing is draining the queue.]",
    "[fleet_submit: could not queue the goal: OSError: disk full]",
    "[fleet_submit: refused -- the goal is 5000 characters, over the 4000 limit.]"])
def test_queue_submitter_other_refusals_are_failures_with_the_message_kept(msg):
    r = QueueSubmitter(fn=lambda *a, **k: msg).submit("[A1] x。")
    assert not r.ok and not r.duplicate and msg[:30] in r.output


def test_queue_submitter_never_raises_from_the_door():
    def boom(*a, **k):
        raise RuntimeError("router import failed")
    r = QueueSubmitter(fn=boom).submit("[A1] x。")
    assert not r.ok and "router import failed" in r.output


def test_queue_submitter_refuses_a_multiline_goal():
    with pytest.raises(ValueError):
        QueueSubmitter(fn=lambda *a, **k: "").submit("a\nb")


def test_queue_submitter_against_the_real_door_in_a_temp_queue(tmp_path, monkeypatch):
    from relay import task_router as TR
    monkeypatch.setattr(TR, "TASKS", str(tmp_path / "tasks"))
    TR.ensure_dirs()
    q = QueueSubmitter()
    goal = "[A1] 先月の件数を数えて。結果の表と図は /tmp/o/A1/ に保存して、要点を3行で答えて。"
    r = q.submit(goal)
    assert r.ok and r.job_id
    files = os.listdir(tmp_path / "tasks" / "pending")
    assert files == [r.job_id + ".json"]
    job = json.loads((tmp_path / "tasks" / "pending" / files[0]).read_text(encoding="utf-8"))
    assert job["payload"]["goal"] == goal and job["origin"]["source"] == "overnight batch"
    assert q.submit(goal).duplicate                         # the door's own duplicate rule


def test_job_waiting_sees_the_queue_directories(tmp_path):
    (tmp_path / "tasks" / "pending").mkdir(parents=True)
    (tmp_path / "tasks" / "pending" / "abc123.json").write_text("{}", encoding="utf-8")
    rd = FleetReader(tmp_path)
    assert rd.job_waiting("abc123") and not rd.job_waiting("zzz")


# ------------------------------------------------------------------ runner behaviour
class _Clock:
    def __init__(self):
        self.t = datetime(2026, 1, 10, 23, 0, tzinfo=JST)
        self.slept = 0.0

    def now(self):
        return self.t

    def sleep(self, s):
        self.slept += s
        self.t += timedelta(seconds=s)


class _Sub:
    def __init__(self, results):
        self.results = list(results)

    def submit(self, goal, probe=None):
        return self.results.pop(0)

    def readonly_check(self):
        return "ro"


class _Reader:
    def __init__(self, seq=None, waiting=False):
        self.seq = seq or {}
        self.waiting = waiting

    def observe(self, items, since, now):
        return {i: (self.seq[i].pop(0) if len(self.seq.get(i, [])) > 1 else self.seq.get(i, [Observation()])[0])
                for i in items}

    def job_waiting(self, jid):
        return self.waiting


def _mk(tmp_path, n, **cfg):
    st = S.State(tmp_path / "state.json")
    st.config.update({"out_dir": str(tmp_path), "preamble": "", "max_minutes": 20, "poll_s": 60,
                      "max_inflight": 10, "submit_route": "queue", **cfg})
    for i in range(1, n + 1):
        st.add_item("A%d" % i, "依頼その%d" % i)
    st.save()
    return st


def test_all_submissions_failing_is_not_reported_as_finished(tmp_path):
    st = _mk(tmp_path, 3)
    fails = [SubmitResult(ok=False, attempts=4, output="screen locked")] * 3
    Runner(st, clock=_Clock(), reader=_Reader(), submitter=_Sub(fails)).run()
    assert [st.get(i)["category"] for i in st.ids()] == [classify.NOT_SUBMITTABLE] * 3
    assert "finished" not in st.config["stop_reason"] and "3 of 3" in st.config["stop_reason"]


def test_a_queued_job_that_the_fleet_has_not_picked_up_is_not_failed_as_unseen(tmp_path):
    st = _mk(tmp_path, 1, unseen_grace_min=5, max_minutes=60)
    clock = _Clock()
    deadline = clock.now() + timedelta(minutes=30)
    st.config["deadline"] = deadline.isoformat()
    ok = [SubmitResult(ok=True, attempts=1, job_id="abc123abc123")]
    Runner(st, clock=clock, reader=_Reader(waiting=True), submitter=_Sub(ok)).run()
    rec = st.get("A1")
    assert rec["job_id"] == "abc123abc123"
    assert "never showed" not in rec["error"]               # waited in the queue, then the deadline


def test_without_the_queue_marker_an_unseen_goal_still_fails(tmp_path):
    st = _mk(tmp_path, 1, unseen_grace_min=5)
    ok = [SubmitResult(ok=True, attempts=1, job_id="abc123abc123")]
    Runner(st, clock=_Clock(), reader=_Reader(waiting=False), submitter=_Sub(ok)).run()
    assert st.get("A1")["status"] == S.FAILED and st.get("A1")["category"] == classify.NOT_SUBMITTABLE


def test_duplicate_rejection_is_reworded_once_on_the_queue_route(tmp_path):
    st = _mk(tmp_path, 1)
    seq = [SubmitResult(duplicate=True), SubmitResult(ok=True, attempts=1, job_id="j1j1j1j1j1j1")]
    done = Observation(state=DONE, outcome="DONE", answer="要点\n1\n2\n3")
    Runner(st, clock=_Clock(), reader=_Reader({"A1": [done]}), submitter=_Sub(seq)).run()
    assert st.get("A1")["variant"] == 1 and st.get("A1")["status"] == S.DONE


# ------------------------------------------------------------------ report banner / reconcile
def _report_state(tmp_path, statuses, route="queue", reconcile_result=None):
    st = S.State(tmp_path / "state.json")
    st.config.update({"out_dir": str(tmp_path), "run_started": "2026-01-10 23:00:00",
                      "run_ended": "2026-01-11 06:00:00", "stop_reason": "no item succeeded",
                      "submit_route": route})
    if reconcile_result is not None:
        st.config["reconcile"] = reconcile_result
    for i, (status, cat) in enumerate(statuses, 1):
        st.add_item("A%d" % i, "依頼その%d" % i)
        st.get("A%d" % i).update({"status": status, "category": cat, "submitted_ts": 1.0,
                                  "error": "e"})
    st.save()
    return st


def test_banner_names_the_count_and_the_likely_cause(tmp_path):
    st = _report_state(tmp_path, [(S.FAILED, classify.NOT_SUBMITTABLE)] * 3 + [(S.DONE, "")], route="gui")
    report.finish(st, tmp_path)
    wb = load_workbook(next(tmp_path.glob("朝レポート_*.xlsx")))
    ws = wb["サマリ"]
    assert "4 件中 3 件が投入できませんでした" in ws["A1"].value and "画面ロック" in ws["A1"].value
    assert ws["A2"].value == "ID" and ws.max_row == 6
    md = next(tmp_path.glob("朝レポート_*.md")).read_text(encoding="utf-8")
    h = next(tmp_path.glob("朝レポート_*.html")).read_text(encoding="utf-8")
    assert "3 件が投入できませんでした" in md and "3 件が投入できませんでした" in h
    fails = wb["失敗と未完"]
    assert fails["D2"].value == classify.NOT_SUBMITTABLE and "再投入" in fails["F2"].value


def test_no_banner_when_everything_was_submitted(tmp_path):
    st = _report_state(tmp_path, [(S.DONE, ""), (S.FAILED, classify.OTHER)])
    report.finish(st, tmp_path)
    ws = load_workbook(next(tmp_path.glob("朝レポート_*.xlsx")))["サマリ"]
    assert ws["A1"].value == "ID"


def test_feedback_still_reads_a_report_that_has_a_banner(tmp_path):
    from tools.overnight import feedback
    run = tmp_path / "r1"
    run.mkdir()
    st = _report_state(run, [(S.FAILED, classify.NOT_SUBMITTABLE), (S.DONE, "")])
    report.finish(st, run)
    x = next(run.glob("朝レポート_*.xlsx"))
    wb = load_workbook(x)
    wb["サマリ"]["I4"].value = "○"
    wb.save(x)
    assert feedback.ingest(x) == 1


def test_reconcile_sheet_and_sections_are_always_present(tmp_path):
    st = _report_state(tmp_path, [(S.DONE, "")])
    report.finish(st, tmp_path)                              # no check was ever run
    wb = load_workbook(next(tmp_path.glob("朝レポート_*.xlsx")))
    assert "整合検査" in wb.sheetnames
    assert "実行されませんでした" in wb["整合検査"]["A2"].value
    (tmp_path / "x").mkdir()
    st2 = _report_state(tmp_path / "x", [(S.DONE, "")],
                        reconcile_result={"checked": 1, "mismatches": ["A1: ずれの例"], "readonly": "rc=0 ok"})
    report.finish(st2, tmp_path / "x")
    md = next((tmp_path / "x").glob("朝レポート_*.md")).read_text(encoding="utf-8")
    assert "## 整合検査" in md and "ずれの例" in md and "rc=0 ok" in md


def test_compare_finds_each_kind_of_disagreement():
    st = S.State("unused.json")
    for iid, status, sub in (("A1", S.DONE, 1.0), ("A2", S.TIMEOUT, 1.0), ("A3", S.DONE, 1.0),
                             ("A4", S.PENDING, None), ("A5", S.DONE, 1.0), ("A6", S.FAILED, 1.0)):
        st.add_item(iid, "q")
        st.get(iid).update({"status": status, "submitted_ts": sub})
    obs = {"A1": Observation(state=DONE), "A2": Observation(state=DONE), "A3": Observation(state=RUNNING),
           "A4": Observation(state=QUEUED), "A5": Observation(state=UNSEEN), "A6": Observation(state=RUNNING)}
    out = reconcile.compare(st, obs)
    ids = [s.split(":")[0] for s in out]
    assert ids == ["A2", "A3", "A4", "A5", "A6"]            # A1 agrees
    assert "完了" in out[0] and "投入していない" in out[2]


def test_run_check_survives_a_failing_readonly_call_and_reader(tmp_path):
    st = _report_state(tmp_path, [(S.DONE, "")])

    class BadReader:
        def observe(self, *a):
            raise OSError("status.json locked")

    def locked():
        raise RuntimeError("cockpit not reachable")

    res = reconcile.run_check(st, BadReader(), locked, 100.0)
    assert "取得できませんでした" in res["readonly"] and "失敗しました" in res["error"]
    assert any("突合した項目数" in l for l in reconcile.describe(res))


def test_run_check_with_matching_workers_reports_no_gap(tmp_path):
    st = _report_state(tmp_path, [(S.DONE, "")])
    (tmp_path / "status.json").write_text(json.dumps({"workers": [
        {"name": "w0", "goal": "[A1] x", "status": "done", "outcome": "DONE"}]}), encoding="utf-8")
    res = reconcile.run_check(st, FleetReader(tmp_path), lambda: "rc=0 fields", 100.0)
    assert res["mismatches"] == [] and "ずれ: なし" in reconcile.describe(res)


# ------------------------------------------------------------------ cli route
def test_cli_defaults_to_the_queue_route_and_gui_can_be_chosen(tmp_path, monkeypatch):
    seen = {}

    class FakeQueue:
        def __init__(self, **k):
            seen["queue"] = True

        def submit(self, goal, probe=None):
            return SubmitResult(ok=True, attempts=1, job_id="q1q1q1q1q1q1")

        def readonly_check(self):
            return "queue"

    class FakeGui(FakeQueue):
        def __init__(self, **k):
            seen["gui"] = True

        def submit(self, goal, probe=None):
            return SubmitResult(ok=True, attempts=1)

    monkeypatch.setattr(cli, "QueueSubmitter", FakeQueue)
    monkeypatch.setattr(cli, "Submitter", FakeGui)
    monkeypatch.setattr(cli.env, "toast", lambda *a, **k: None)
    monkeypatch.setattr(cli.env, "free_disk_gb", lambda p: 100.0)
    monkeypatch.setattr(cli.env, "disk_floor_gb", lambda: 1.0)
    fleet = tmp_path / "fleet"
    fleet.mkdir()
    items = tmp_path / "items.txt"
    items.write_text("件数を数えて\n", encoding="utf-8")
    base = ["run", "--items", str(items), "--fleet-dir", str(fleet), "--poll-seconds", "0",
            "--max-minutes-per-item", "0.01"]
    assert cli.batch_main(base + ["--out", str(tmp_path / "q")]) == 0
    st = S.State.load(tmp_path / "q" / "state.json")
    assert st.config["submit_route"] == "queue" and st.get("I001")["job_id"] == "q1q1q1q1q1q1"
    assert (tmp_path / "q" / "READY.txt").exists()
    assert st.config["reconcile"]["checked"] == 1            # the morning check ran
    seen.clear()
    assert cli.batch_main(base + ["--out", str(tmp_path / "g"), "--submit-route", "gui"]) == 0
    assert S.State.load(tmp_path / "g" / "state.json").config["submit_route"] == "gui"


# ------------------------------------------------------------------ similar goals are held back
def test_similar_requests_are_serialised_instead_of_refused(tmp_path):
    from tools.overnight.submit import fleet_similarity
    sim = fleet_similarity()
    assert sim is not None
    st = _mk(tmp_path, 3)
    calls = []

    class Sub:
        def submit(self, goal, probe=None):
            calls.append(goal)
            return SubmitResult(ok=True, attempts=1, job_id="j%011d" % len(calls))

        def readonly_check(self):
            return "ro"

    done = lambda: Observation(state=DONE, outcome="DONE", answer="要点\n1\n2\n3")
    # nothing finishes until a few polls have passed: the second goal must not be sent early
    r = _Reader({"A1": [Observation(state=RUNNING), done()], "A2": [done()], "A3": [done()]})
    Runner(st, clock=_Clock(), reader=r, submitter=Sub(), similar=sim).run()
    assert len(calls) == 3
    assert [st.get(i)["status"] for i in st.ids()] == [S.DONE] * 3


def test_a_request_is_not_held_back_when_nothing_is_similar(tmp_path):
    st = _mk(tmp_path, 2)
    sent = []

    class Sub:
        def submit(self, goal, probe=None):
            sent.append(goal)
            return SubmitResult(ok=True, attempts=1)

        def readonly_check(self):
            return "ro"

    class R(_Reader):
        def observe(self, items, since, now):
            return {i: Observation(state=RUNNING) for i in items}

    clock = _Clock()
    st.config["deadline"] = (clock.now() + timedelta(minutes=3)).isoformat()
    Runner(st, clock=clock, reader=R(), submitter=Sub(), similar=lambda a, b: False).run()
    assert len(sent) == 2                                   # both went out in the first sweep
