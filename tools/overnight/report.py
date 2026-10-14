"""The morning report: one xlsx (editable), the same content as md and as self-contained html.

Whatever happened overnight, finish() leaves something readable: a failed xlsx falls back to md,
a failed md falls back to plain text, and every error goes to report_error.txt. READY.txt is
written last and says what exists.
"""
from __future__ import annotations

import csv
import html
import re
import statistics
import traceback
from pathlib import Path

from tools.overnight import classify, reconcile
from tools.overnight import state as S
from tools.overnight.common import Clock, fmt_dt, from_epoch

SUMMARY = "サマリ"
FAILS = "失敗と未完"
INFO = "実行情報"
RECONCILE = "整合検査"
RESULT_LABEL = {S.DONE: "成功", S.FAILED: "失敗", S.TIMEOUT: "時間切れ"}
UNRUN = "未実施"
CELL_MAX = 30000
TABLE_MAX_ROWS = 200
TABLE_MAX_COLS = 20
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SQL_FENCE = re.compile(r"```(?:sql|SQL)?\s*\n(.*?)```", re.S)
_SQL_START = re.compile(r"^\s*(select|with)\b", re.I)
_INTERP = re.compile(r"解釈|前提|仮定|とみなし|として扱|と読み替|曖昧")
_LIST_PREFIX = re.compile(r"^\s*(?:[-*・•]|\d+[.)、]|[①-⑩])\s*")


def clean(text) -> str:
    t = ILLEGAL.sub("", "" if text is None else str(text))
    return t[:CELL_MAX]


# ------------------------------------------------------------------ view model
def result_label(rec) -> str:
    return RESULT_LABEL.get(rec["status"], UNRUN)


def key_points(answer: str, n: int = 3) -> list:
    lines = []
    for raw in (answer or "").replace("```", "\n").splitlines():
        s = raw.strip()
        if s and s.upper() != "DONE":
            lines.append(s)
    idx = next((i for i, s in enumerate(lines) if "要点" in s), None)
    picked = []
    if idx is not None:
        tail = lines[idx + 1:idx + 1 + n]
        head = re.sub(r"^.*要点[^:：]*[:：]?\s*", "", lines[idx]).strip("*# ")
        picked = ([head] if head else []) + tail
        picked = picked[:n]
    if not picked:
        picked = lines[-n:]
    return [_LIST_PREFIX.sub("", p).strip("*# ")[:160] for p in picked if p]


def extract_sql(answer: str) -> list:
    found = []
    for m in _SQL_FENCE.finditer(answer or ""):
        body = m.group(1).strip()
        if _SQL_START.match(body):
            found.append(body)
    if not found:
        buf = []
        for ln in (answer or "").splitlines():
            if _SQL_START.match(ln) or (buf and ln.strip()):
                buf.append(ln.rstrip())
            elif buf:
                found.append("\n".join(buf))
                buf = []
        if buf:
            found.append("\n".join(buf))
    return found[:5]


def interpretation_notes(answer: str) -> list:
    out = []
    for ln in (answer or "").splitlines():
        if _INTERP.search(ln):
            out.append(ln.strip()[:240])
    return out[:5]


def list_artifacts(out_dir: Path, iid: str) -> list:
    d = Path(out_dir) / iid
    files = []
    try:
        for p in sorted(d.rglob("*")):
            if p.is_file() and p.name not in ("_answer.txt",):
                files.append(p.relative_to(out_dir).as_posix())
            if len(files) >= 50:
                break
    except OSError:
        pass
    return files


def duration_min(rec):
    start = rec.get("started_ts") or rec.get("submitted_ts")
    end = rec.get("finished_ts")
    if not start or not end or end < start:
        return None
    return round((end - start) / 60.0, 1)


def read_table(path: Path):
    """First sheet of an xlsx, or a csv -> list of rows (capped); [] when unreadable."""
    rows = []
    try:
        if path.suffix.lower() == ".csv":
            raw = path.read_bytes()
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp932", errors="replace")
            for i, r in enumerate(csv.reader(text.splitlines())):
                if i >= TABLE_MAX_ROWS:
                    break
                rows.append([_num_or_text(c) for c in r[:TABLE_MAX_COLS]])
        else:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.worksheets[0]
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= TABLE_MAX_ROWS:
                    break
                rows.append(list(r[:TABLE_MAX_COLS]))
            wb.close()
    except Exception:
        return []
    return rows


def _num_or_text(c):
    try:
        return int(c)
    except (TypeError, ValueError):
        try:
            return float(c)
        except (TypeError, ValueError):
            return c


def build_view(st: S.State) -> dict:
    out_dir = Path(st.config["out_dir"])
    rows = []
    for iid in st.ids():
        rec = st.get(iid)
        lab = result_label(rec)
        ans = rec.get("answer") or ""
        cat = rec.get("category") or ""
        failed = lab in ("失敗", "時間切れ", UNRUN)
        files = list_artifacts(out_dir, iid)
        rows.append({
            "id": iid, "ask": rec["ask"], "label": lab, "status": rec["status"],
            "points": key_points(ans) if ans else [],
            "minutes": duration_min(rec), "files": files,
            "category": cat if failed else "",
            "error": rec.get("error") or "", "note": rec.get("note") or "",
            "next": ("実行されませんでした。python -m tools.overnight resume --out <この出力フォルダ> --retry-failed で続きを実行できます。"
                     if rec["status"] == S.SKIPPED else
                     classify.next_step_for(cat or classify.OTHER)) if failed else "",
            "sql": extract_sql(ans), "interp": interpretation_notes(ans),
            "answer": ans, "outcome": rec.get("outcome") or "", "turns": rec.get("turns"),
            "worker": rec.get("worker") or "",
        })
    done = sum(1 for r in rows if r["status"] == S.DONE)
    mins = [r["minutes"] for r in rows if r["minutes"] is not None and r["status"] != S.SKIPPED]
    counts = {k: sum(1 for r in rows if r["label"] == k) for k in ("成功", "失敗", "時間切れ", UNRUN)}
    dist = {}
    if mins:
        dist = {"最短": min(mins), "中央値": round(statistics.median(mins), 1), "最長": max(mins),
                "5分未満": sum(m < 5 for m in mins), "5-10分": sum(5 <= m < 10 for m in mins),
                "10-20分": sum(10 <= m < 20 for m in mins), "20分以上": sum(m >= 20 for m in mins)}
    cfg = st.config
    info = [
        ("開始", cfg.get("run_started", "")), ("終了", cfg.get("run_ended", "")),
        ("終了理由", cfg.get("stop_reason", "")), ("項目数", len(rows)),
        ("成功", counts["成功"]), ("失敗", counts["失敗"]), ("時間切れ", counts["時間切れ"]),
        ("未実施", counts[UNRUN]),
        ("成功率", "%.0f%%" % (100.0 * done / len(rows)) if rows else "-"),
        ("項目ごとの時間予算(分)", cfg.get("max_minutes", "")),
        ("同時投入の上限", cfg.get("max_inflight", "")),
        ("期限(--deadline)", cfg.get("deadline", "") or "なし"),
        ("開始指定(--start-at)", cfg.get("start_at", "") or "なし"),
    ]
    for k, v in dist.items():
        info.append(("所要時間 " + k, v))
    for k, v in (cfg.get("fleet_summary") or {}).items():
        info.append(("フリート " + str(k), v))
    for w in cfg.get("warnings") or []:
        info.append(("入力の注意", w))
    end = cfg.get("run_ended") or fmt_dt(Clock().now())
    info.insert(3, ("投入経路", {"queue": "キュー(fleet_submit)", "gui": "コックピット画面操作"}.get(
        cfg.get("submit_route", ""), cfg.get("submit_route", "") or "不明")))
    blocked = sum(1 for r in rows if r["category"] == classify.NOT_SUBMITTABLE)
    banner = ""
    if blocked:
        cause = ("画面ロック、コックピットの未起動" if cfg.get("submit_route") == "gui"
                 else "キューの満杯・拒否、フリートの停止") if cfg.get("submit_route") else \
            "画面ロック(GUI経路)、またはキューの満杯・拒否"
        banner = "%d 件中 %d 件が投入できませんでした。原因候補: %s。" % (len(rows), blocked, cause)
    return {"rows": rows, "info": info, "date": end[:10].replace("-", ""), "counts": counts,
            "banner": banner, "stop_reason": cfg.get("stop_reason", ""),
            "reconcile": reconcile.describe(cfg.get("reconcile"))}


# ------------------------------------------------------------------ xlsx
def _chart_for(ws, table_rows, top_row: int, anchor_col: int):
    """Native Excel bar chart over the first text column and first numeric column of the table.

    Deliberately small and self-contained so a shared chart tool can replace it later.
    """
    from openpyxl.chart import BarChart, Reference
    from openpyxl.utils import get_column_letter
    if len(table_rows) < 3:
        return False
    header, body = table_rows[0], table_rows[1:]
    num_col = next((c for c in range(len(header))
                    if all(isinstance(r[c], (int, float)) for r in body[:20] if len(r) > c and r[c] is not None)
                    and any(isinstance(r[c], (int, float)) for r in body if len(r) > c)), None)
    if num_col is None:
        return False
    cat_col = next((c for c in range(len(header)) if c != num_col), None)
    if cat_col is None:
        return False
    last = top_row + min(len(table_rows) - 1, 30)
    ch = BarChart()
    ch.type = "col"
    ch.title = clean(header[num_col])
    ch.height, ch.width = 8, 16
    ch.add_data(Reference(ws, min_col=num_col + 1, min_row=top_row, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col + 1, min_row=top_row + 1, max_row=last))
    ws.add_chart(ch, "%s%d" % (get_column_letter(anchor_col), top_row))
    return True


def _sheet_name(iid: str, used: set) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", "_", iid)[:31] or "item"
    name, n = base, 2
    while name in used or name in (SUMMARY, FAILS, INFO, RECONCILE):
        suffix = "_%d" % n
        name = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def write_xlsx(view: dict, out_dir: Path, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    fills = {"成功": "E2EFDA", "失敗": "F8CBAD", "時間切れ": "FFE699", UNRUN: "D9D9D9"}
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, cols, widths):
        ws.append(cols)
        hdr = ws.max_row
        for i, c in enumerate(ws[hdr], 1):
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
            ws.column_dimensions[c.column_letter].width = widths[i - 1]
        ws.freeze_panes = "A%d" % (hdr + 1)
        return hdr

    wb = Workbook()
    ws = wb.active
    ws.title = SUMMARY
    if view.get("banner"):
        ws.append([clean(view["banner"])])
        ws.merge_cells("A1:J1")
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="C00000")
        ws["A1"].alignment = wrap
        ws.row_dimensions[1].height = 32
    hdr_row = header(ws, ["ID", "依頼文", "結果", "要点(3行)", "所要時間(分)", "成果物", "失敗の分類", "詳細",
                "OK/NG", "SQL(参考)"], [8, 40, 10, 60, 12, 28, 18, 8, 8, 30])
    used: set = set()
    names = {}
    for r in view["rows"]:
        names[r["id"]] = _sheet_name(r["id"], used)
    for r in view["rows"]:
        ws.append([clean(r["id"]), clean(r["ask"]), r["label"], clean("\n".join(r["points"])),
                   r["minutes"] if r["minutes"] is not None else "", "",
                   clean(r["category"]), "開く", "", clean("\n\n".join(r["sql"]))])
        row = ws.max_row
        for c in ws[row]:
            c.alignment = wrap
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=fills.get(r["label"], "FFFFFF"))
        if r["files"]:
            first = r["files"][0]
            cell = ws.cell(row, 6, clean(first))
            cell.hyperlink = first
            cell.font = Font(color="0563C1", underline="single")
        elif (Path(out_dir) / r["id"]).is_dir():
            cell = ws.cell(row, 6, "フォルダ")
            cell.hyperlink = r["id"] + "/"
            cell.font = Font(color="0563C1", underline="single")
        link = ws.cell(row, 8)
        link.hyperlink = "#'%s'!A1" % names[r["id"]]
        link.font = Font(color="0563C1", underline="single")
    if view["rows"]:
        dv = DataValidation(type="list", formula1='"○,×"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("I%d:I%d" % (hdr_row + 1, hdr_row + len(view["rows"])))

    wf = wb.create_sheet(FAILS)
    header(wf, ["ID", "依頼文", "状態", "原因の分類", "詳細", "翌朝にやること"], [8, 40, 10, 18, 50, 50])
    bad = [r for r in view["rows"] if r["label"] != "成功"]
    for r in bad:
        wf.append([clean(r["id"]), clean(r["ask"]), r["label"], clean(r["category"] or "-"),
                   clean(r["error"]), clean(r["next"] or "-")])
    for r in view["rows"]:
        if r["label"] == "成功" and r["note"]:
            wf.append([clean(r["id"]), clean(r["ask"]), "成功(要確認)", "", clean(r["note"]),
                       "回答を読んで、成功と見てよいか確認してください。"])
    if wf.max_row == 1:
        wf.append(["", "失敗・未完の項目はありません。", "", "", "", ""])
    for row in wf.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    wi = wb.create_sheet(INFO)
    header(wi, ["項目", "値"], [30, 60])
    for k, v in view["info"]:
        wi.append([clean(k), clean(v)])

    wr = wb.create_sheet(RECONCILE)
    header(wr, ["整合検査(レポートとフリートの突き合わせ)"], [120])
    for line in view["reconcile"]:
        wr.append([clean(line)])
        wr.cell(wr.max_row, 1).alignment = wrap

    for r in view["rows"]:
        wsi = wb.create_sheet(names[r["id"]])
        wsi.column_dimensions["A"].width = 18
        wsi.column_dimensions["B"].width = 100
        meta = [("ID", r["id"]), ("依頼文", r["ask"]), ("結果", r["label"]),
                ("要点", "\n".join(r["points"])), ("所要時間(分)", r["minutes"] or ""),
                ("失敗の分類", r["category"]), ("エラー", r["error"]), ("注記", r["note"]),
                ("使用SQL", "\n\n".join(r["sql"]) or "(回答に記載なし)"),
                ("解釈の注記", "\n".join(r["interp"]) or "(記載なし)"),
                ("成果物", "\n".join(r["files"]) or "(なし)"),
                ("作業者の回答", r["answer"])]
        for k, v in meta:
            wsi.append([k, clean(v)])
            wsi.cell(wsi.max_row, 1).font = Font(bold=True)
            wsi.cell(wsi.max_row, 1).alignment = wrap
            wsi.cell(wsi.max_row, 2).alignment = wrap
        tbl = None
        for f in sorted(r["files"], key=lambda n: not n.lower().endswith(".xlsx")):
            if f.lower().endswith((".xlsx", ".csv")) and not Path(f).name.startswith("~$"):
                tbl = read_table(Path(out_dir) / f)
                if tbl:
                    break
        if tbl:
            top = wsi.max_row + 3
            wsi.cell(top - 1, 1, "取り込んだ表").font = Font(bold=True)
            for i, row_vals in enumerate(tbl):
                for j, v in enumerate(row_vals):
                    wsi.cell(top + i, 1 + j, clean(v) if isinstance(v, str) else v)
            try:
                _chart_for(wsi, tbl, top, len(tbl[0]) + 2)
            except Exception:
                pass
        pngs = [f for f in r["files"] if f.lower().endswith(".png")][:2]
        for k, f in enumerate(pngs):
            try:
                from openpyxl.drawing.image import Image
                img = Image(str(Path(out_dir) / f))
                scale = min(1.0, 640.0 / max(1, img.width))
                img.width, img.height = int(img.width * scale), int(img.height * scale)
                wsi.add_image(img, "D%d" % (2 + k * 25))
            except Exception:
                pass
    wb.save(path)


# ------------------------------------------------------------------ md / html
def _md_cell(t) -> str:
    return str(t).replace("|", "\\|").replace("\n", "<br>")


def write_md(view: dict, path: Path) -> None:
    L = ["# 朝レポート %s" % view["date"], ""]
    c = view["counts"]
    L.append("成功 %d / 失敗 %d / 時間切れ %d / 未実施 %d" % (c["成功"], c["失敗"], c["時間切れ"], c[UNRUN]))
    if view.get("banner"):
        L[1:1] = ["", "> **" + view["banner"] + "**"]
    L += ["", "## サマリ", "", "| ID | 依頼文 | 結果 | 要点 | 分 | 成果物 | 失敗の分類 | OK/NG |",
          "|---|---|---|---|---|---|---|---|"]
    for r in view["rows"]:
        links = ", ".join("[%s](%s)" % (Path(f).name, f.replace(" ", "%20")) for f in r["files"][:5])
        L.append("| %s | %s | %s | %s | %s | %s | %s | |" % (
            r["id"], _md_cell(r["ask"]), r["label"], _md_cell("\n".join(r["points"])),
            "" if r["minutes"] is None else r["minutes"], _md_cell(links), r["category"]))
    L += ["", "## 失敗と未完", ""]
    bad = [r for r in view["rows"] if r["label"] != "成功"]
    if not bad:
        L.append("なし")
    for r in bad:
        L.append("- **%s** %s (%s%s): %s -> %s" % (
            r["id"], r["ask"], r["label"], "/" + r["category"] if r["category"] else "",
            r["error"][:200], r["next"]))
    L += ["", "## 実行情報", ""]
    L += ["- %s: %s" % (k, v) for k, v in view["info"]]
    L += ["", "## 整合検査", ""] + ["- " + x for x in view["reconcile"]]
    L += ["", "## 項目ごと", ""]
    for r in view["rows"]:
        L += ["### %s %s (%s)" % (r["id"], r["ask"], r["label"]), ""]
        L += ["- " + p for p in r["points"]] or ["(回答なし)"]
        if r["sql"]:
            L += ["", "使用SQL:", "", "```sql", "\n\n".join(r["sql"]), "```"]
        if r["interp"]:
            L += ["", "解釈の注記:"] + ["- " + t for t in r["interp"]]
        if r["note"]:
            L += ["", "注記: " + r["note"]]
        L.append("")
    path.write_text("\n".join(L), encoding="utf-8")


_CSS = ("body{font-family:'Yu Gothic','Meiryo',sans-serif;margin:24px;color:#222}"
        "table{border-collapse:collapse;width:100%}th,td{border:1px solid #bbb;padding:4px 8px;"
        "vertical-align:top;font-size:13px}th{background:#1F4E78;color:#fff}"
        ".成功{background:#E2EFDA}.失敗{background:#F8CBAD}.時間切れ{background:#FFE699}"
        ".未実施{background:#D9D9D9}pre{background:#f4f4f4;padding:8px;overflow:auto}"
        "details{margin:8px 0}")


def write_html(view: dict, path: Path) -> None:
    e = html.escape
    c = view["counts"]
    P = ["<!doctype html><html lang='ja'><head><meta charset='utf-8'>",
         "<meta name='viewport' content='width=device-width,initial-scale=1'>",
         "<title>朝レポート %s</title><style>%s</style></head><body>" % (e(view["date"]), _CSS),
         "<h1>朝レポート %s</h1>" % e(view["date"]),
         "<p>成功 %d / 失敗 %d / 時間切れ %d / 未実施 %d</p>" % (c["成功"], c["失敗"], c["時間切れ"], c[UNRUN]),
         ("<p style='background:#C00000;color:#fff;padding:8px;font-weight:bold'>%s</p>" % e(view["banner"])) if view.get("banner") else "",
         "<h2>サマリ</h2><table><tr><th>ID</th><th>依頼文</th><th>結果</th><th>要点</th><th>分</th>"
         "<th>成果物</th><th>失敗の分類</th></tr>"]
    for r in view["rows"]:
        links = " ".join("<a href='%s'>%s</a>" % (e(f.replace(" ", "%20")), e(Path(f).name))
                         for f in r["files"][:5])
        P.append("<tr><td><a href='#i-%s'>%s</a></td><td>%s</td><td class='%s'>%s</td><td>%s</td>"
                 "<td>%s</td><td>%s</td><td>%s</td></tr>" % (
                     e(r["id"]), e(r["id"]), e(r["ask"]), r["label"], r["label"],
                     "<br>".join(e(p) for p in r["points"]),
                     "" if r["minutes"] is None else r["minutes"], links, e(r["category"])))
    P.append("</table><h2>失敗と未完</h2>")
    bad = [r for r in view["rows"] if r["label"] != "成功"]
    if not bad:
        P.append("<p>なし</p>")
    else:
        P.append("<ul>")
        for r in bad:
            P.append("<li><b>%s</b> %s (%s %s): %s<br>%s</li>" % (
                e(r["id"]), e(r["ask"]), r["label"], e(r["category"]), e(r["error"][:200]), e(r["next"])))
        P.append("</ul>")
    P.append("<h2>実行情報</h2><table>")
    P += ["<tr><th>%s</th><td>%s</td></tr>" % (e(str(k)), e(str(v))) for k, v in view["info"]]
    P.append("</table><h2>整合検査</h2><ul>%s</ul><h2>項目ごと</h2>" % "".join("<li>%s</li>" % e(x) for x in view["reconcile"]))
    for r in view["rows"]:
        P.append("<details id='i-%s'><summary>%s %s (%s)</summary>" % (
            e(r["id"]), e(r["id"]), e(r["ask"]), r["label"]))
        P.append("<ul>%s</ul>" % "".join("<li>%s</li>" % e(p) for p in r["points"]))
        if r["sql"]:
            P.append("<p>使用SQL</p><pre>%s</pre>" % e("\n\n".join(r["sql"])))
        if r["interp"]:
            P.append("<p>解釈の注記</p><ul>%s</ul>" % "".join("<li>%s</li>" % e(t) for t in r["interp"]))
        if r["note"]:
            P.append("<p>注記: %s</p>" % e(r["note"]))
        P.append("</details>")
    P.append("</body></html>")
    path.write_text("\n".join(P), encoding="utf-8")


# ------------------------------------------------------------------ entry
def _plain_fallback(st: S.State, path: Path) -> None:
    lines = ["朝レポート(簡易版: 通常の作成に失敗したため)", ""]
    for iid in st.ids():
        rec = st.get(iid)
        lines.append("%s | %s | %s | %s" % (iid, rec["status"], rec["ask"], (rec.get("error") or "")[:120]))
    path.write_text("\n".join(lines), encoding="utf-8")


def finish(st: S.State, out_dir, log=None, toast=None, clock=None) -> dict:
    """Write the report files, READY.txt, and (best effort) a toast. Never raises."""
    log = log or (lambda m: None)
    out = Path(out_dir)
    made, errors = [], []
    view = None
    try:
        out.mkdir(parents=True, exist_ok=True)
        view = build_view(st)
    except Exception:
        errors.append("build_view:\n" + traceback.format_exc())
    date = view["date"] if view else (clock or Clock()).now().strftime("%Y%m%d")
    targets = [("xlsx", lambda p: write_xlsx(view, out, p)), ("md", lambda p: write_md(view, p)),
               ("html", lambda p: write_html(view, p))] if view else []
    for ext, fn in targets:
        p = out / ("朝レポート_%s.%s" % (date, ext))
        try:
            fn(p)
            made.append(p.name)
        except Exception:
            errors.append("%s:\n%s" % (ext, traceback.format_exc()))
            try:
                if p.exists() and ext == "xlsx":
                    p.unlink()
            except OSError:
                pass
    if not any(n.endswith(".md") for n in made):
        p = out / ("朝レポート_%s.md" % date)
        try:
            _plain_fallback(st, p)
            made.append(p.name)
        except Exception:
            errors.append("plain fallback:\n" + traceback.format_exc())
    if errors:
        try:
            (out / "report_error.txt").write_text("\n\n".join(errors), encoding="utf-8")
        except OSError:
            pass
        log("report problems were written to report_error.txt")
    counts = view["counts"] if view else {}
    try:
        (out / "READY.txt").write_text(
            "morning report ready\nfiles: %s\nerrors: %d\nresults: %s\n"
            % (", ".join(made) or "(none)", len(errors), counts), encoding="utf-8")
    except OSError:
        pass
    if toast is not None:
        try:
            toast("朝レポートができました", "成功 %s / 全 %s 件。%s" % (
                counts.get("成功", "?"), len(st.ids()), out))
        except Exception:
            pass
    log("report done: %s" % ", ".join(made))
    return {"files": made, "errors": len(errors)}
