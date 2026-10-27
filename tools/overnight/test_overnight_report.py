# -*- coding: utf-8 -*-
"""Report generation (xlsx / md / html), failure-proofing, feedback, and the CLI wiring."""
from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from tools.overnight import batch_cli as cli, feedback, report
from tools.overnight import state as S
from tools.overnight.common import JST
from tools.overnight.submit import SubmitResult
from tools.overnight.submit import SubmitResult

ANSWER = ("調べました。\n解釈: 「先月」は直近の暦月と解釈しました。\n```sql\nSELECT a, COUNT(*) FROM t GROUP BY a\n```\n"
          "要点\n- 件数は10件\n- 最多はX\n- 増加傾向\nDONE")


def _state(tmp_path, spec):
    st = S.State(tmp_path / "state.json")
    st.config.update({"out_dir": str(tmp_path), "run_started": "2026-01-10 23:00:00",
                      "run_ended": "2026-01-11 06:30:00", "stop_reason": "finished",
                      "max_minutes": 20, "fleet_summary": {"max_concurrent": 7}, "warnings": ["w1"]})
    for i, (status, answer, cat) in enumerate(spec, 1):
        iid = "A%d" % i
        st.add_item(iid, "依頼その%d" % i)
        rec = st.get(iid)
        rec.update({"status": status, "answer": answer, "category": cat,
                    "started_ts": 1000.0, "submitted_ts": 900.0, "finished_ts": 1000.0 + 60 * (i + 1),
                    "error": "e%d" % i})
    st.save()
    return st


def _worker_files(tmp_path):
    d = tmp_path / "A1"
    d.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["品種", "件数"])
    for n, v in (("a", 3), ("b", 5), ("c", 2)):
        ws.append([n, v])
    wb.save(d / "result.xlsx")
    (d / "extra.csv").write_text("x,y\n1,2\n", encoding="utf-8")


def test_report_has_all_sheets_links_failures_and_chart(tmp_path):
    st = _state(tmp_path, [(S.DONE, ANSWER, ""), (S.FAILED, "DBに接続できません", "DB接続不可"),
                           (S.TIMEOUT, "", "タイムアウト"), (S.SKIPPED, "", "")])
    _worker_files(tmp_path)
    res = report.finish(st, tmp_path)
    assert res["errors"] == 0 and len(res["files"]) == 3
    x = next(tmp_path.glob("朝レポート_*.xlsx"))
    wb = load_workbook(x)
    assert wb.sheetnames[:3] == ["サマリ", "失敗と未完", "実行情報"]
    assert {"A1", "A2", "A3", "A4"} <= set(wb.sheetnames)
    ws = wb["サマリ"]
    assert ws.max_row == 5
    assert [ws.cell(1, c).value for c in (1, 3, 9)] == ["ID", "結果", "OK/NG"]
    assert [ws.cell(r, 3).value for r in range(2, 6)] == ["成功", "失敗", "時間切れ", "未実施"]
    assert ws["F2"].hyperlink is not None and ws["F2"].hyperlink.target.startswith("A1/")
    assert ws["H2"].hyperlink is not None
    assert "件数は10件" in ws["D2"].value and "SELECT" in ws["J2"].value
    fails = wb["失敗と未完"]
    assert fails.max_row == 4 and fails["D2"].value == "DB接続不可"
    assert "再実行" in fails["F2"].value or "確認" in fails["F2"].value
    info = {wb["実行情報"].cell(r, 1).value: wb["実行情報"].cell(r, 2).value
            for r in range(2, wb["実行情報"].max_row + 1)}
    assert info["項目数"] == "4" and info["成功率"] == "25%" and info["フリート max_concurrent"] == "7"
    item = wb["A1"]
    assert len(item._charts) == 1                       # native Excel chart over the imported table
    labels = [item.cell(r, 1).value for r in range(1, 14)]
    assert "使用SQL" in labels and "解釈の注記" in labels
    flat = " ".join(str(c.value) for row in item.iter_rows() for c in row if c.value is not None)
    assert "暦月と解釈" in flat and "品種" in flat


def test_md_and_html_are_self_contained_and_escaped(tmp_path):
    st = _state(tmp_path, [(S.DONE, ANSWER + "\n<script>alert(1)</script>", ""), (S.FAILED, "x", "その他")])
    report.finish(st, tmp_path)
    md = next(tmp_path.glob("朝レポート_*.md")).read_text(encoding="utf-8")
    h = next(tmp_path.glob("朝レポート_*.html")).read_text(encoding="utf-8")
    assert "## サマリ" in md and "## 失敗と未完" in md and "SELECT a" in md
    assert "<script" not in h.lower().replace("<script>alert", "") or "&lt;script" in h
    assert "src=" not in h and "http://" not in h and "https://" not in h
    assert "依頼その2" in h


def test_no_success_still_produces_a_report(tmp_path):
    st = _state(tmp_path, [(S.FAILED, "x", "その他"), (S.TIMEOUT, "", "タイムアウト")])
    res = report.finish(st, tmp_path)
    assert len(res["files"]) == 3 and (tmp_path / "READY.txt").exists()
    wb = load_workbook(next(tmp_path.glob("朝レポート_*.xlsx")))
    assert wb["サマリ"].max_row == 3


def test_empty_state_still_produces_a_report(tmp_path):
    st = _state(tmp_path, [])
    assert len(report.finish(st, tmp_path)["files"]) == 3


def test_report_failure_leaves_report_error_and_a_plain_report(tmp_path, monkeypatch):
    st = _state(tmp_path, [(S.DONE, ANSWER, "")])

    def boom(*a, **k):
        raise RuntimeError("xlsx exploded")

    monkeypatch.setattr(report, "write_xlsx", boom)
    monkeypatch.setattr(report, "write_html", boom)
    res = report.finish(st, tmp_path)
    err = (tmp_path / "report_error.txt").read_text(encoding="utf-8")
    assert res["errors"] == 2 and "xlsx exploded" in err
    assert next(tmp_path.glob("朝レポート_*.md")).exists() and (tmp_path / "READY.txt").exists()


def test_total_failure_falls_back_to_plain_text(tmp_path, monkeypatch):
    st = _state(tmp_path, [(S.DONE, ANSWER, "")])
    monkeypatch.setattr(report, "build_view", lambda s: (_ for _ in ()).throw(ValueError("bad view")))
    res = report.finish(st, tmp_path)
    assert "bad view" in (tmp_path / "report_error.txt").read_text(encoding="utf-8")
    md = next(tmp_path.glob("朝レポート_*.md")).read_text(encoding="utf-8")
    assert "依頼その1" in md and (tmp_path / "READY.txt").exists() and res["errors"] == 1


def test_toast_failure_is_ignored(tmp_path):
    st = _state(tmp_path, [(S.DONE, ANSWER, "")])

    def bad_toast(*a):
        raise OSError("no shell")

    assert report.finish(st, tmp_path, toast=bad_toast)["errors"] == 0


def test_control_characters_and_long_text_do_not_break_the_sheet(tmp_path):
    st = _state(tmp_path, [(S.DONE, "x\x00y\x07" + "あ" * 50000, "")])
    assert report.finish(st, tmp_path)["errors"] == 0


def test_sheet_names_are_valid_and_unique(tmp_path):
    st = S.State(tmp_path / "state.json")
    st.config.update({"out_dir": str(tmp_path)})
    for iid in ("a/b", "a:b", "x" * 40, "x" * 40 + "y", "サマリ"):
        st.add_item(iid, "q")
    st.save()
    assert report.finish(st, tmp_path)["errors"] == 0
    names = load_workbook(next(tmp_path.glob("朝レポート_*.xlsx"))).sheetnames
    assert len(names) == len(set(names)) == 9 and all(len(n) <= 31 for n in names)


def test_key_points_sql_and_interpretation_extraction():
    assert report.key_points(ANSWER) == ["件数は10件", "最多はX", "増加傾向"]
    assert report.key_points("a\nb\nc\nd\nDONE") == ["b", "c", "d"]
    assert report.extract_sql(ANSWER) == ["SELECT a, COUNT(*) FROM t GROUP BY a"]
    assert report.extract_sql("SELECT 1\nFROM t\n\nok") == ["SELECT 1\nFROM t"]
    assert report.interpretation_notes(ANSWER)[0].startswith("解釈")


# ------------------------------------------------------------------ feedback
def _marked_report(tmp_path, marks):
    st = _state(tmp_path, [(S.DONE, ANSWER, ""), (S.DONE, ANSWER, ""), (S.DONE, ANSWER, "")])
    report.finish(st, tmp_path)
    x = next(tmp_path.glob("朝レポート_*.xlsx"))
    wb = load_workbook(x)
    for i, m in enumerate(marks, 2):
        wb["サマリ"].cell(i, 9).value = m
    wb.save(x)
    return x


def test_feedback_reads_marks_and_is_idempotent(tmp_path):
    run = tmp_path / "20260110_2300"
    run.mkdir()
    x = _marked_report(run, ["○", "×", None])
    assert feedback.ingest(x) == 2
    assert feedback.ingest(x) == 0
    lines = [json.loads(l) for l in (tmp_path / "feedback.jsonl").read_text(encoding="utf-8").splitlines()]
    assert [(l["id"], l["verdict"]) for l in lines] == [("A1", "ok"), ("A2", "ng")]
    assert "SELECT" in lines[0]["sql"] and lines[0]["ask"] == "依頼その1"


def test_feedback_marks_variants():
    v = feedback.verdict_of
    assert [v(x) for x in ("○", "OK", " ng ", "×", "?", None, "")] == ["ok", "ok", "ng", "ng", "", "", ""]


def test_feedback_picks_up_earlier_nights_at_the_next_run(tmp_path):
    prev = tmp_path / "20260109_2300"
    prev.mkdir()
    _marked_report(prev, ["○", None, None])
    now = tmp_path / "20260110_2300"
    now.mkdir()
    assert feedback.ingest_siblings(now) == 1
    assert feedback.ingest_siblings(now) == 0


def test_feedback_on_a_broken_file_is_reported_not_raised(tmp_path):
    bad = tmp_path / "r" / "朝レポート_1.xlsx"
    bad.parent.mkdir()
    bad.write_bytes(b"not an xlsx")
    assert feedback.ingest_siblings(tmp_path / "other") == 0
    assert cli.batch_main(["feedback", "--report", str(bad)]) == 2


# ------------------------------------------------------------------ cli
def test_cli_run_end_to_end_with_a_fake_fleet_and_resume(tmp_path, monkeypatch):
    items = tmp_path / "items.txt"
    items.write_text("# 前置き: 簡潔に\n件数を数えて\n平均を出して\n件数を数えて\n", encoding="utf-8")
    fleet = tmp_path / "fleet"
    fleet.mkdir()
    out = tmp_path / "out" / "n1"
    sent = []

    class FakeSub:
        def __init__(self, **k):
            pass

        def submit(self, goal, probe=None):
            sent.append(goal)
            return SubmitResult(ok=True, attempts=1)

        def readonly_check(self):
            return "ro"

    def write_status():
        ws = [{"name": "w%d" % i, "goal": g, "status": "done", "outcome": "DONE", "last": ANSWER,
               "phase_events": [{"ts": 1.0, "event": "waiting"}]} for i, g in enumerate(sent)]
        (fleet / "status.json").write_text(json.dumps({"running": True, "workers": ws}), encoding="utf-8")

    class Sub2(FakeSub):
        def submit(self, goal, probe=None):
            r = super().submit(goal, probe)
            write_status()
            return r

    monkeypatch.setattr(cli, "Submitter", Sub2)
    monkeypatch.setattr(cli.env, "toast", lambda *a, **k: None)
    monkeypatch.setattr(cli.env, "free_disk_gb", lambda p: 100.0)
    monkeypatch.setattr(cli.env, "disk_floor_gb", lambda: 1.0)
    rc = cli.batch_main(["run", "--items", str(items), "--out", str(out), "--fleet-dir", str(fleet), "--submit-route", "gui",
                   "--poll-seconds", "0"])
    assert rc == 0
    # All THREE items are submitted. The two identical goal lines ("件数を数えて")
    # are no longer folded into one: the owner's 2026-10-26/27 order (fix/no-content-based-skip)
    # is that a content look-alike must never stop a job -- only a duplicate item ID is rejected
    # (tools/overnight/items.py). Each still arrives rewritten to a single bracketed [Innn] line.
    assert len(sent) == 3 and all("\n" not in g and g.startswith("[") for g in sent)
    assert (out / "READY.txt").exists() and (out / "overnight.log").exists()
    wb = load_workbook(next(out.glob("朝レポート_*.xlsx")))
    rows = [(wb["サマリ"].cell(r, 1).value, wb["サマリ"].cell(r, 3).value) for r in range(2, wb["サマリ"].max_row + 1)]
    # All three ran and succeeded -- the duplicate goal line is no longer reported "未実施"
    # (unsubmitted), because it is no longer folded out; it is submitted as I003 like the rest.
    assert ("I001", "成功") in rows and ("I002", "成功") in rows and ("I003", "成功") in rows
    # running it again over the same folder is refused; resume sends nothing new
    assert cli.batch_main(["run", "--items", str(items), "--out", str(out)]) == 2
    n = len(sent)
    assert cli.batch_main(["resume", "--out", str(out), "--fleet-dir", str(fleet), "--poll-seconds", "0"]) == 0
    assert len(sent) == n


def test_cli_rejects_a_file_with_nothing_to_run(tmp_path):
    p = tmp_path / "items.txt"
    p.write_text("# nothing\n", encoding="utf-8")
    assert cli.batch_main(["run", "--items", str(p), "--out", str(tmp_path / "o")]) == 2
